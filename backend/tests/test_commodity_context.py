from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import Workbook
from fastapi.testclient import TestClient

import trendforge_api.main as main_module
from trendforge_api.commodity_context import (
    build_commodity_context_snapshot,
    latest_commodity_context_snapshot,
    save_commodity_context_snapshot,
)
from trendforge_api.institutional_sources import EndpointFetchResult, FetchState
from trendforge_api.main import app


def _result(
    endpoint_key: str,
    payload,
    *,
    url: str,
    raw_path: str | None = None,
    state: FetchState = FetchState.RAW_ARCHIVED,
) -> EndpointFetchResult:
    return EndpointFetchResult(
        endpointKey=endpoint_key,
        state=state,
        fetchedAt=datetime(2026, 7, 14, 12, tzinfo=UTC),
        url=url,
        contentHash=f"hash-{endpoint_key}",
        rawPath=raw_path,
        mediaType="application/json",
        recordCount=len(payload.get("Data", [])) if isinstance(payload, dict) else 1,
        payload=payload,
        canScore=False,
    )


def test_mcx_option_chain_calculates_pcr_walls_and_true_max_pain() -> None:
    payload = {
        "Summary": {"AsOn": "/Date(1784024699000)/", "Count": 3},
        "Data": [
            {
                "CE_StrikePrice": 100,
                "CE_OpenInterest": 100,
                "CE_ChangeInOI": 10,
                "CE_Volume": 20,
                "CE_LTP": 12,
                "PE_OpenInterest": 20,
                "PE_ChangeInOI": 2,
                "PE_Volume": 3,
                "PE_LTP": 1,
                "UnderlyingValue": 110,
            },
            {
                "CE_StrikePrice": 110,
                "CE_OpenInterest": 40,
                "CE_ChangeInOI": 4,
                "CE_Volume": 8,
                "CE_LTP": 6,
                "PE_OpenInterest": 200,
                "PE_ChangeInOI": 20,
                "PE_Volume": 30,
                "PE_LTP": 5,
                "UnderlyingValue": 110,
            },
            {
                "CE_StrikePrice": 120,
                "CE_OpenInterest": 20,
                "CE_ChangeInOI": -2,
                "CE_Volume": 5,
                "CE_LTP": 2,
                "PE_OpenInterest": 50,
                "PE_ChangeInOI": 5,
                "PE_Volume": 7,
                "PE_LTP": 10,
                "UnderlyingValue": 110,
            },
        ],
    }
    result = _result(
        "mcx_option_chain",
        payload,
        url=(
            "https://www.mcxindia.com/GetOptionChain?"
            "InstrumentType=optfut&Symbol=GOLD&Expiry=29JUL2026"
        ),
    )

    snapshot = build_commodity_context_snapshot([result])

    metrics = snapshot.mcx_option_metrics[0]
    assert metrics.symbol == "GOLD"
    assert metrics.expiry == "29JUL2026"
    assert metrics.call_open_interest == 160
    assert metrics.put_open_interest == 270
    assert metrics.pcr_oi == 1.6875
    assert metrics.call_wall_strike == 100
    assert metrics.put_wall_strike == 110
    assert metrics.max_pain_strike == 110
    assert metrics.signal_state == "CONTEXT_AVAILABLE"
    assert len(snapshot.mcx_option_rows) == 6
    assert snapshot.can_unlock_ready is False


def test_mcx_zero_open_interest_is_unknown_not_bullish_or_bearish() -> None:
    result = _result(
        "mcx_option_chain",
        {
            "Summary": {"AsOn": "/Date(1784024699000)/", "Count": 1},
            "Data": [
                {
                    "CE_StrikePrice": 100,
                    "CE_OpenInterest": 0,
                    "PE_OpenInterest": 0,
                    "UnderlyingValue": 100,
                }
            ],
        },
        url=(
            "https://www.mcxindia.com/GetOptionChain?"
            "InstrumentType=optfut&Symbol=GOLD&Expiry=29JUL2026"
        ),
    )

    metrics = build_commodity_context_snapshot([result]).mcx_option_metrics[0]

    assert metrics.pcr_oi is None
    assert metrics.max_pain_strike is None
    assert metrics.call_wall_strike is None
    assert metrics.put_wall_strike is None
    assert metrics.signal_state == "WAIT_NO_OPEN_INTEREST"


def test_mcx_market_watch_normalizes_contract_quotes_as_research_context() -> None:
    result = _result(
        "mcx_market_watch",
        {
            "success": True,
            "data": {
                "Data": [
                    {
                        "Symbol": "CRUDEOIL",
                        "ProductCode": "CRUDEOIL",
                        "ExpiryDate": "20JUL2026",
                        "InstrumentName": "FUTCOM",
                        "Open": 7_720,
                        "High": 7_810,
                        "Low": 7_690,
                        "LTP": 7_795,
                        "PreviousClose": 7_715,
                        "PercentChange": 1.04,
                        "Volume": 12_500,
                        "OpenInterest": 18_200,
                        "LTTValue": "2026-07-14 15:57:57",
                        "BuyPrice": 7_794,
                        "SellPrice": 7_796,
                        "BuyQuantity": 10,
                        "SellQuantity": 8,
                    }
                ]
            },
        },
        url="https://www.mcxindia.com/market-data/market-watch/GetMarketWatch?culture=en",
    )

    snapshot = build_commodity_context_snapshot([result])

    assert snapshot.sources[0].parser_state == "STRUCTURED_RESEARCH_ONLY"
    assert len(snapshot.mcx_market_rows) == 1
    quote = snapshot.mcx_market_rows[0]
    assert quote.symbol == "CRUDEOIL"
    assert quote.expiry == "20JUL2026"
    assert quote.data_date == "2026-07-14"
    assert quote.last_price == 7_795
    assert quote.open_interest == 18_200
    assert quote.bid_ask_spread == 2
    assert quote.can_unlock_ready is False


def test_sge_is_benchmark_trend_only_and_does_not_invent_lbma_premium() -> None:
    result = _result(
        "sge_benchmark_gold",
        {
            "zp": [[1783728000000, 750.0], [1783814400000, 752.0]],
            "wp": [[1783728000000, 751.0], [1783814400000, 753.5]],
        },
        url="https://en.sge.com.cn/graph/DayilyJzj",
    )

    metrics = build_commodity_context_snapshot([result]).sge_benchmark

    assert metrics is not None
    assert metrics.latest_zp == 752.0
    assert metrics.latest_wp == 753.5
    assert metrics.intraday_spread_proxy == 1.5
    assert metrics.lbma_premium is None
    assert "LBMA" in metrics.limitation


def test_baker_hughes_workbook_normalizes_supply_context(tmp_path: Path) -> None:
    path = tmp_path / "rig-count.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "NAM Summary"
    sheet["D3"] = "NORTH AMERICA Rotary Rig Count"
    sheet["D4"] = "10/07/2026"
    sheet["B19"] = "U.S. Breakout Information"
    for row_number, values in {
        11: ("United States Total", 581, 1, 580, 44, 537),
        15: ("Canada", 179, -11, 190, 17, 162),
        17: ("North America", 760, -10, 770, 61, 699),
        21: ("Gas", 126, 0, 126, 18, 108),
        22: ("Oil", 445, 0, 445, 21, 424),
    }.items():
        sheet.cell(row_number, 2, values[0])
        sheet.cell(row_number, 4, values[1])
        sheet.cell(row_number, 5, values[2])
        sheet.cell(row_number, 6, values[3])
        sheet.cell(row_number, 7, values[4])
        sheet.cell(row_number, 8, values[5])
    workbook.save(path)
    result = _result(
        "baker_hughes_na_rig_count",
        {"byteLength": path.stat().st_size},
        url="https://bakerhughesrigcount.gcs-web.com/static-files/example",
        raw_path=str(path),
    )

    metrics = build_commodity_context_snapshot([result]).baker_hughes

    assert metrics is not None
    assert metrics.data_date == "2026-07-10"
    assert metrics.us_total == 581
    assert metrics.us_oil == 445
    assert metrics.us_gas == 126
    assert metrics.north_america_weekly_change == -10
    assert metrics.supply_signal == "SUPPLY_CONTRACTION"


def test_context_snapshot_persists_metrics_and_row_lineage(tmp_path: Path) -> None:
    result = _result(
        "mcx_option_chain",
        {
            "Summary": {"AsOn": "/Date(1784024699000)/", "Count": 1},
            "Data": [
                {
                    "CE_StrikePrice": 100,
                    "CE_OpenInterest": 10,
                    "PE_OpenInterest": 20,
                    "UnderlyingValue": 100,
                }
            ],
        },
        url=(
            "https://www.mcxindia.com/GetOptionChain?"
            "InstrumentType=optfut&Symbol=GOLD&Expiry=29JUL2026"
        ),
    )
    snapshot = build_commodity_context_snapshot([result])
    database = tmp_path / "research.sqlite3"

    save_commodity_context_snapshot(snapshot, db_path=database)
    restored = latest_commodity_context_snapshot(db_path=database)

    assert restored.run_id == snapshot.run_id
    assert restored.mcx_option_metrics[0].pcr_oi == 2.0
    assert restored.mcx_option_rows[0].source_content_hash == "hash-mcx_option_chain"
    assert json.loads(restored.model_dump_json())["canUnlockReady"] is False


def test_latest_commodity_context_api_is_read_only_and_uses_aliases(monkeypatch) -> None:
    snapshot = build_commodity_context_snapshot(
        [
            _result(
                "mcx_option_chain",
                {
                    "Summary": {"AsOn": "/Date(1784024699000)/", "Count": 1},
                    "Data": [
                        {
                            "CE_StrikePrice": 100,
                            "CE_OpenInterest": 10,
                            "PE_OpenInterest": 20,
                            "UnderlyingValue": 100,
                        }
                    ],
                },
                url=(
                    "https://www.mcxindia.com/GetOptionChain?"
                    "InstrumentType=optfut&Symbol=GOLD&Expiry=29JUL2026"
                ),
            )
        ]
    )
    monkeypatch.setattr(
        main_module,
        "latest_commodity_context_snapshot",
        lambda: snapshot,
    )

    response = TestClient(app).get("/api/institutional/commodity-context/latest")

    assert response.status_code == 200
    payload = response.json()
    assert payload["state"] == "RESEARCH_ONLY"
    assert payload["canUnlockReady"] is False
    assert payload["mcxOptionMetrics"][0]["pcrOi"] == 2.0
