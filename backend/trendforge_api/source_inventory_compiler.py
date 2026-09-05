"""R0 inventory compiler — Hybrid §16.3–16.6 / §18.4 / §19.2 detail under File A ceilings.

Research-only: builds inventory defect reports, decision-job maps, dataset-root
identity fields, and maturity-stage counts. Does not fetch market data, place
orders, or claim GATE_AUTHORIZED from REGISTERED/CONNECTED labels alone.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter
from dataclasses import dataclass
from datetime import datetime, timezone
from enum import StrEnum
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
from urllib.parse import urlparse

from pydantic import BaseModel, ConfigDict, Field

from .source_monitor import SOURCE_CATALOG
from .source_key_map_review import (
    EndpointSourceMapRecord,
    SourceKeyMapRecord,
    build_source_key_map,
)
from .source_overlap_resolutions import (
    REVIEWED_OVERLAP_RESOLUTIONS,
    OverlapResolution,
    resolve_dataset_root_identity,
    REVIEWED_WORKBOOK_SHA256,
    OverlapDisposition,
    OverlapResolutionState,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK = (
    PROJECT_ROOT / "data" / "reports" / "SOURCE_LINK_INVENTORY_MASTER.xlsx"
)

# Hybrid §16.4 known sentinel tokens that must not live in canonical URL fields.
KNOWN_SENTINEL_URLS = frozenset(
    {
        "lme_warehouse_stocks:no_snapshot",
        "mca_master_data:no_snapshot",
        "mcx_delivery_reports:no_snapshot",
        "rbi_fpi_monitoring:no_snapshot",
    }
)

JSON_ARRAY_KEYS = re.compile(r"^\s*\[")
HTTP_URL = re.compile(r"https?://[^|\s]+", re.IGNORECASE)


def _to_camel(value: str) -> str:
    first, *rest = value.split("_")
    return first + "".join(part.capitalize() for part in rest)


class DecisionJob(StrEnum):
    """Hybrid §19.2 J01–J14. Jobs never unlock execution under File A research scope."""

    J01 = "J01"
    J02 = "J02"
    J03 = "J03"
    J04 = "J04"
    J05 = "J05"
    J06 = "J06"
    J07 = "J07"
    J08 = "J08"
    J09 = "J09"
    J10 = "J10"
    J11 = "J11"
    J12 = "J12"
    J13 = "J13"
    J14 = "J14"


class CompilerState(StrEnum):
    OK = "OK"
    INPUT_MISSING = "INPUT_MISSING"
    INPUT_EMPTY = "INPUT_EMPTY"
    SCHEMA_ERROR = "SCHEMA_ERROR"


DECISION_JOB_SPECS: dict[DecisionJob, dict[str, str]] = {
    DecisionJob.J01: {
        "purpose": "Identity, symbols, ISINs, calendars and contract masters",
        "scope": "No directional vote",
    },
    DecisionJob.J02: {
        "purpose": "Tradability restrictions, bans, bands, halts, auctions and surveillance",
        "scope": "Veto or research WAIT/REJECT only; no quantity UI",
    },
    DecisionJob.J03: {
        "purpose": "Price, volume, delivery, spread, depth and liquidity",
        "scope": "Timeframe-specific market evidence",
    },
    DecisionJob.J04: {
        "purpose": "Pre-open auction state and open confirmation",
        "scope": "Session-bound; previous close cannot substitute",
    },
    DecisionJob.J05: {
        "purpose": "Market regime, breadth, index and sector leadership",
        "scope": "Context; not symbol-specific institutional proof",
    },
    DecisionJob.J06: {
        "purpose": "Futures OI, basis, rollover and participant positioning",
        "scope": "Participant data remains regime context",
    },
    DecisionJob.J07: {
        "purpose": "Options, IV, Greeks, PCR, Max Pain, OI walls and GEX proxy",
        "scope": "One parent-chain validity ceiling",
    },
    DecisionJob.J08: {
        "purpose": "Deals, PIT, SAST, pledge and ownership events",
        "scope": "Event identity and direction required",
    },
    DecisionJob.J09: {
        "purpose": "Corporate filings, results, announcements and corporate actions",
        "scope": "Publication-time and revision aware",
    },
    DecisionJob.J10: {
        "purpose": "AMFI and NSDL institutional context",
        "scope": "Delayed swing evidence only",
    },
    DecisionJob.J11: {
        "purpose": "MCX local price, OI, contracts, expiry, margin and delivery state",
        "scope": "Local MCX evidence required for MCX research confirm path",
    },
    DecisionJob.J12: {
        "purpose": "Global commodity positioning, inventory, physical and macro context",
        "scope": "Context cannot replace local MCX evidence",
    },
    DecisionJob.J13: {
        "purpose": "Portfolio risk, costs, sizing and stressed realizable loss",
        "scope": "POSTPONE product quantity; may only produce research warnings",
    },
    DecisionJob.J14: {
        "purpose": "Outcomes, calibration, drift and model degradation",
        "scope": "Point-in-time research only",
    },
}


class SourceMaturityState(StrEnum):
    """Hybrid §16.3 ladder. EXECUTION_AUTHORIZED is out of current product scope."""

    REGISTERED = "REGISTERED"
    TRANSPORT_OK = "TRANSPORT_OK"
    ARTIFACT_VALID = "ARTIFACT_VALID"
    VALID_EMPTY = "VALID_EMPTY"
    PARSER_SCHEMA_OK = "PARSER_SCHEMA_OK"
    NORMALIZED = "NORMALIZED"
    FRESH_FOR_JOB = "FRESH_FOR_JOB"
    DECISION_WIRED = "DECISION_WIRED"
    GATE_AUTHORIZED = "GATE_AUTHORIZED"
    EXECUTION_AUTHORIZED = "EXECUTION_AUTHORIZED"


MATURITY_LADDER: tuple[SourceMaturityState, ...] = tuple(SourceMaturityState)

# Extended source-contract fields (File A AMEND-A-001 / Hybrid §15.3).
EXTENDED_SOURCE_CONTRACT_FIELDS: tuple[str, ...] = (
    "timezone",
    "publication_calendar",
    "session_dependency",
    "entitlement_required",
    "legal_use_mode",
    "rate_budget",
    "retry_policy",
    "circuit_breaker_policy",
    "content_signature",
    "schema_version",
    "correction_policy",
    "event_identity_fields",
    "fallback_authority",
    "watermark_policy",
    "retention_policy",
)

# Dataset-root identity fields (File A AMEND-A-005 / Hybrid §18.4).
DATASET_ROOT_IDENTITY_FIELDS: tuple[str, ...] = (
    "dataset_root_id",
    "transport_variant_id",
    "mirror_group_id",
    "resolver_id",
    "publisher_authority",
    "endpoint_role",
    "activation_state",
    "decision_jobs",
    "evidence_family",
    "family_weight_cap",
    "fallback_authority_cap",
)


class InventoryDefectCode(StrEnum):
    EXACT_URL_DUPLICATE = "EXACT_URL_DUPLICATE"
    KEY_PREFIXED_URL = "KEY_PREFIXED_URL"
    COMPOUND_URL = "COMPOUND_URL"
    SENTINEL_URL = "SENTINEL_URL"
    MISSING_PURPOSE_JOBS = "MISSING_PURPOSE_JOBS"
    PROSE_PURPOSE_JOBS = "PROSE_PURPOSE_JOBS"
    JSON_ARRAY_ACTIVE_KEYS = "JSON_ARRAY_ACTIVE_KEYS"
    PLACEHOLDER_OR_TEST = "PLACEHOLDER_OR_TEST"
    MISSING_ROLE = "MISSING_ROLE"
    NON_HTTP_CANONICAL = "NON_HTTP_CANONICAL"

    NON_DATA_REFERENCE_ROW = "NON_DATA_REFERENCE_ROW"
    SEMANTIC_SOURCE_OVERLAP = "SEMANTIC_SOURCE_OVERLAP"
    REGISTRY_RUNTIME_CONTRACT_MISMATCH = "REGISTRY_RUNTIME_CONTRACT_MISMATCH"


# File B section 16.4 defects 1-11. Extra compiler-only defects may coexist,
# but this list is the stable TDG-GAP-012 contract exposed by the API.
HYBRID_H1A0_DEFECT_CODES: tuple[InventoryDefectCode, ...] = (
    InventoryDefectCode.EXACT_URL_DUPLICATE,
    InventoryDefectCode.KEY_PREFIXED_URL,
    InventoryDefectCode.COMPOUND_URL,
    InventoryDefectCode.SENTINEL_URL,
    InventoryDefectCode.MISSING_PURPOSE_JOBS,
    InventoryDefectCode.PROSE_PURPOSE_JOBS,
    InventoryDefectCode.PLACEHOLDER_OR_TEST,
    InventoryDefectCode.NON_DATA_REFERENCE_ROW,
    InventoryDefectCode.JSON_ARRAY_ACTIVE_KEYS,
    InventoryDefectCode.SEMANTIC_SOURCE_OVERLAP,
    InventoryDefectCode.REGISTRY_RUNTIME_CONTRACT_MISMATCH,
)


class CompilerRow(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    inventory_id: str = Field(alias="inventoryId")
    url: str
    source_role: str | None = Field(default=None, alias="sourceRole")
    purpose_jobs: str | None = Field(default=None, alias="purposeJobs")
    active_source_keys: str | None = Field(default=None, alias="activeSourceKeys")
    connection_status: str | None = Field(default=None, alias="connectionStatus")
    parser_status: str | None = Field(default=None, alias="parserStatus")
    freshness_status: str | None = Field(default=None, alias="freshnessStatus")
    current_usability: str | None = Field(default=None, alias="currentUsability")
    placeholder_or_test: str | None = Field(default=None, alias="placeholderOrTest")
    sheet: str = "MASTER_CURRENT"


class InventoryDefect(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    code: InventoryDefectCode
    inventory_id: str = Field(alias="inventoryId")
    url: str
    detail: str


class NormalizedEndpointRecord(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    endpoint_id: str = Field(alias="endpointId")
    canonical_url: str = Field(alias="canonicalUrl")
    source_contract_ids: list[str] = Field(alias="sourceContractIds")
    inventory_ids: list[str] = Field(alias="inventoryIds")
    partition_memberships: list[str] = Field(alias="partitionMemberships")
    endpoint_role: str = Field(alias="endpointRole")
    transport_variant_id: str = Field(alias="transportVariantId")
    publisher_authority: str = Field(alias="publisherAuthority")
    decision_jobs: list[str] = Field(alias="decisionJobs")
    safe_use: str = Field(alias="safeUse")
    blocked_use: str = Field(alias="blockedUse")
    next_action: str = Field(alias="nextAction")
    quarantined: bool = False
    quarantine_reasons: list[str] = Field(
        default_factory=list, alias="quarantineReasons"
    )

    purpose: str


class MigrationAction(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    action_id: str = Field(alias="actionId")
    inventory_id: str = Field(alias="inventoryId")
    action_type: str = Field(alias="actionType")
    original_value: str = Field(alias="originalValue")
    normalized_urls: list[str] = Field(alias="normalizedUrls")
    source_contract_ids: list[str] = Field(alias="sourceContractIds")
    reason: str
    destructive: bool = False


class CompiledSourceContract(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    source_contract_id: str = Field(alias="sourceContractId")
    dataset_root_id: str = Field(alias="datasetRootId")
    endpoint_ids: list[str] = Field(alias="endpointIds")
    inventory_ids: list[str] = Field(alias="inventoryIds")
    artifact_kind: str = Field(alias="artifactKind")
    parser_id: str | None = Field(default=None, alias="parserId")
    parser_version: str | None = Field(default=None, alias="parserVersion")
    dataset_id: str | None = Field(default=None, alias="datasetId")
    decision_jobs: list[str] = Field(alias="decisionJobs")
    feature_ids: list[str] = Field(default_factory=list, alias="featureIds")
    strategy_profiles: list[str] = Field(default_factory=list, alias="strategyProfiles")
    mandatory_for: list[str] = Field(default_factory=list, alias="mandatoryFor")
    confirmation_for: list[str] = Field(default_factory=list, alias="confirmationFor")
    veto_for: list[str] = Field(default_factory=list, alias="vetoFor")
    independence_family: str = Field(alias="independenceFamily")
    authority_cap: str = Field(alias="authorityCap")
    freshness_by_job: dict[str, str] = Field(alias="freshnessByJob")
    fallback_chain: list[str] = Field(default_factory=list, alias="fallbackChain")
    primary_panel_field: str | None = Field(default=None, alias="primaryPanelField")
    inspector_sections: list[str] = Field(
        default_factory=list, alias="inspectorSections"
    )
    tests: list[str] = Field(default_factory=list)
    allowed_timeframes: list[str] = Field(
        default_factory=list, alias="allowedTimeframes"
    )
    allowed_instruments: list[str] = Field(
        default_factory=list, alias="allowedInstruments"
    )
    directional_permission: bool = Field(default=False, alias="directionalPermission")
    veto_permission: bool = Field(default=False, alias="vetoPermission")
    quantity_permission: bool = Field(default=False, alias="quantityPermission")
    gate_permission: bool = Field(default=False, alias="gatePermission")
    maturity_state: SourceMaturityState = Field(alias="maturityState")
    maturity_missing_prerequisites: list[str] = Field(
        default_factory=list, alias="maturityMissingPrerequisites"
    )
    extended_fields: dict[str, str] = Field(alias="extendedFields")
    missing_extended_fields: list[str] = Field(alias="missingExtendedFields")
    contract_status: str = Field(alias="contractStatus")
    safe_use: str = Field(alias="safeUse")
    blocked_use: str = Field(alias="blockedUse")
    next_action: str = Field(alias="nextAction")
    linked: bool = False


class DatasetRootRecord(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    dataset_root_id: str = Field(alias="datasetRootId")
    transport_variant_id: str = Field(alias="transportVariantId")
    mirror_group_id: str = Field(alias="mirrorGroupId")
    resolver_id: str = Field(alias="resolverId")
    publisher_authority: str = Field(alias="publisherAuthority")
    endpoint_role: str = Field(alias="endpointRole")
    activation_state: str = Field(alias="activationState")
    decision_jobs: list[str] = Field(alias="decisionJobs")
    evidence_family: str = Field(alias="evidenceFamily")
    family_weight_cap: float = Field(alias="familyWeightCap")
    fallback_authority_cap: str = Field(alias="fallbackAuthorityCap")
    source_keys: list[str] = Field(default_factory=list, alias="sourceKeys")
    sample_urls: list[str] = Field(default_factory=list, alias="sampleUrls")
    linked: bool = False
    can_unlock_execution: bool = Field(default=False, alias="canUnlockExecution")


class ContractSplitRecord(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    contract_id: str = Field(alias="contractId")
    role: str
    decision_jobs: list[str] = Field(alias="decisionJobs")
    valid_empty_allowed: bool = Field(alias="validEmptyAllowed")
    failure_wait_reason: str = Field(alias="failureWaitReason")
    notes: str
    runtime_alias: str | None = Field(default=None, alias="runtimeAlias")


class AcceptanceCheck(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    id: str
    description: str
    passed: bool
    evidence: str


class OverlapResolutionRecord(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    canonical_url: str = Field(alias="canonicalUrl")
    contract_ids: list[str] = Field(alias="contractIds")
    disposition: OverlapDisposition | None = None
    state: OverlapResolutionState
    canonical_dataset_root_id: str | None = Field(
        default=None, alias="canonicalDatasetRootId"
    )
    inventory_ids: list[str] = Field(default_factory=list, alias="inventoryIds")
    decision_jobs: list[str] = Field(default_factory=list, alias="decisionJobs")
    evidence_reason: str = Field(alias="evidenceReason")
    evidence_basis: list[str] = Field(default_factory=list, alias="evidenceBasis")
    resolution_version: str | None = Field(default=None, alias="resolutionVersion")
    reviewed_at: str | None = Field(default=None, alias="reviewedAt")
    gate_eligible: bool = Field(default=False, alias="gateEligible")


class InventoryCompilerReport(BaseModel):
    model_config = ConfigDict(
        alias_generator=_to_camel, populate_by_name=True, frozen=True
    )

    compiled_at: datetime = Field(alias="compiledAt")
    ok: bool
    compiler_state: CompilerState = Field(alias="compilerState")
    workbook_path: str | None = Field(default=None, alias="workbookPath")
    workbook_sha256: str | None = Field(default=None, alias="workbookSha256")
    row_count: int = Field(alias="rowCount")
    role_counts: dict[str, int] = Field(alias="roleCounts")
    linked_source_key_count: int = Field(alias="linkedSourceKeyCount")
    not_linked_row_count: int = Field(alias="notLinkedRowCount")
    active_catalog_key_count: int = Field(alias="activeCatalogKeyCount")
    source_key_map: list[SourceKeyMapRecord] = Field(alias="sourceKeyMap")
    endpoint_source_map: list[EndpointSourceMapRecord] = Field(
        alias="endpointSourceMap"
    )
    source_map_review_version: str = Field(alias="sourceMapReviewVersion")
    source_map_review_valid: bool = Field(alias="sourceMapReviewValid")
    source_map_review_errors: list[str] = Field(alias="sourceMapReviewErrors")
    source_governance_ready: bool = Field(alias="sourceGovernanceReady")
    reviewed_source_key_count: int = Field(alias="reviewedSourceKeyCount")
    unreviewed_source_key_count: int = Field(alias="unreviewedSourceKeyCount")
    activation_reviewed_source_key_count: int = Field(
        alias="activationReviewedSourceKeyCount"
    )
    gate_authorized_source_key_count: int = Field(alias="gateAuthorizedSourceKeyCount")
    legacy_unresolved_compiled_identity_count: int = Field(
        alias="legacyUnresolvedCompiledIdentityCount"
    )
    named_inventory_source_key_count: int = Field(alias="namedInventorySourceKeyCount")
    generated_endpoint_lineage_count: int = Field(alias="generatedEndpointLineageCount")
    runtime_catalog_source_key_count: int = Field(alias="runtimeCatalogSourceKeyCount")
    shared_source_key_count: int = Field(alias="sharedSourceKeyCount")
    inventory_only_source_key_count: int = Field(alias="inventoryOnlySourceKeyCount")
    runtime_only_source_key_count: int = Field(alias="runtimeOnlySourceKeyCount")
    living_source_key_count: int = Field(alias="livingSourceKeyCount")
    unassigned_endpoint_lineage_count: int = Field(
        alias="unassignedEndpointLineageCount"
    )
    named_inventory_source_key_sha256: str = Field(
        alias="namedInventorySourceKeySha256"
    )
    runtime_catalog_source_key_sha256: str = Field(
        alias="runtimeCatalogSourceKeySha256"
    )
    living_source_key_sha256: str = Field(alias="livingSourceKeySha256")
    defect_counts: dict[str, int] = Field(alias="defectCounts")
    h1a0_defect_codes: list[str] = Field(alias="h1a0DefectCodes")
    defects: list[InventoryDefect]
    overlap_resolution_records: list[OverlapResolutionRecord] = Field(
        default_factory=list, alias="overlapResolutionRecords"
    )
    overlap_disposition_counts: dict[str, int] = Field(
        default_factory=dict, alias="overlapDispositionCounts"
    )
    invalid_overlap_resolution_count: int = Field(
        default=0, alias="invalidOverlapResolutionCount"
    )
    stale_overlap_resolution_count: int = Field(
        default=0, alias="staleOverlapResolutionCount"
    )
    resolution_registry_valid: bool = Field(
        default=True, alias="resolutionRegistryValid"
    )
    maturity_stage_counts: dict[str, int] = Field(alias="maturityStageCounts")
    decision_jobs: list[dict[str, str]] = Field(alias="decisionJobs")
    dataset_roots: list[DatasetRootRecord] = Field(alias="datasetRoots")
    normalized_endpoints: list[NormalizedEndpointRecord] = Field(
        alias="normalizedEndpoints"
    )
    source_contracts: list[CompiledSourceContract] = Field(alias="sourceContracts")
    migration_manifest: list[MigrationAction] = Field(alias="migrationManifest")
    input_canonical_count: int = Field(alias="inputCanonicalCount")
    normalized_endpoint_count: int = Field(alias="normalizedEndpointCount")
    normalized_source_contract_count: int = Field(alias="normalizedSourceContractCount")
    sentinel_status_count: int = Field(alias="sentinelStatusCount")
    migration_action_count: int = Field(alias="migrationActionCount")
    lineage_row_count: int = Field(alias="lineageRowCount")
    roots_truncated: bool = Field(default=False, alias="rootsTruncated")
    source_activation_ready: bool = Field(default=False, alias="sourceActivationReady")
    unresolved_contract_count: int = Field(alias="unresolvedContractCount")
    contract_splits: list[ContractSplitRecord] = Field(alias="contractSplits")
    extended_source_contract_fields: list[str] = Field(
        alias="extendedSourceContractFields"
    )
    dataset_root_identity_fields: list[str] = Field(alias="datasetRootIdentityFields")
    acceptance: list[AcceptanceCheck]
    h1a0_acceptance_passed: int = Field(alias="h1a0AcceptancePassed")
    h1a0_acceptance_total: int = Field(alias="h1a0AcceptanceTotal")
    research_only: bool = Field(default=True, alias="researchOnly")
    execution_authorized_count: int = Field(default=0, alias="executionAuthorizedCount")
    rule: str


@dataclass(frozen=True)
class _RawInventoryRow:
    inventory_id: str
    url: str
    source_role: str | None
    purpose_jobs: str | None
    active_source_keys: str | None
    connection_status: str | None
    parser_status: str | None
    freshness_status: str | None
    current_usability: str | None
    placeholder_or_test: str | None
    sheet: str
    extra: Mapping[str, Any]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def parse_active_source_keys(raw: str | None) -> list[str]:
    if raw is None:
        return []
    text = str(raw).strip()
    if not text or text.lower() in {"none", "null", "nan"}:
        return []
    if JSON_ARRAY_KEYS.match(text):
        try:
            parsed = json.loads(text.replace("'", '"'))
        except json.JSONDecodeError:
            return [text]
        if isinstance(parsed, list):
            return [str(item).strip() for item in parsed if str(item).strip()]
        return [text]
    parts: list[str] = []
    for token in re.split(r"[|,:;/]+", text):
        cleaned = token.strip().strip("'\"")
        if cleaned:
            parts.append(cleaned)
    return parts


def _list_field(value: Any) -> list[str]:
    if value is None:
        return []
    if isinstance(value, (list, tuple, set)):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    if not text or text.lower() in {"none", "null", "nan"}:
        return []
    if text.startswith("["):
        try:
            parsed = json.loads(text.replace("'", '"'))
        except json.JSONDecodeError:
            parsed = None
        if isinstance(parsed, list):
            return [str(item).strip() for item in parsed if str(item).strip()]
    return [item.strip() for item in re.split(r"[|,;]+", text) if item.strip()]


def normalize_inventory_urls(raw: str) -> list[str]:
    """Return distinct HTTP(S) endpoints from a raw inventory cell."""

    urls: list[str] = []
    for match in HTTP_URL.findall(raw or ""):
        cleaned = match.strip().rstrip("),.;]")
        parsed = urlparse(cleaned)
        if parsed.scheme.lower() in {"http", "https"} and parsed.netloc:
            normalized = parsed._replace(
                scheme=parsed.scheme.lower(), netloc=parsed.netloc.lower()
            ).geturl()
            if normalized not in urls:
                urls.append(normalized)
    return urls


def _stable_id(prefix: str, value: str) -> str:
    return f"{prefix}:{hashlib.sha256(value.encode('utf-8')).hexdigest()[:16]}"


def _source_contract_bindings(
    row: _RawInventoryRow, urls: Sequence[str]
) -> dict[str, list[str]]:
    keys = sorted(set(parse_active_source_keys(row.active_source_keys)))
    if keys:
        return {url: keys for url in urls}

    bindings: dict[str, list[str]] = {}
    for key, raw_url in re.findall(
        r"(?:^|\|)\s*([a-z0-9_]+):(https?://[^|]+)",
        _canonical_input_value(row),
        re.I,
    ):
        normalized = normalize_inventory_urls(raw_url)
        if normalized and normalized[0] in urls:
            bindings.setdefault(normalized[0], []).append(key.lower())
    return {
        url: sorted(set(bindings.get(url) or [_stable_id("source", url)]))
        for url in urls
    }


def _source_contract_ids(row: _RawInventoryRow, urls: Sequence[str]) -> list[str]:
    keys = sorted(set(parse_active_source_keys(row.active_source_keys)))
    if not urls:
        return keys
    return sorted(
        {
            contract_id
            for bound_ids in _source_contract_bindings(row, urls).values()
            for contract_id in bound_ids
        }
    )


def _explicit_or_inferred_jobs(row: _RawInventoryRow, keys: Sequence[str]) -> list[str]:
    explicit = sorted(
        {
            token.upper()
            for token in _list_field(row.purpose_jobs)
            if re.fullmatch(r"J(?:0[1-9]|1[0-4])", token.upper())
        }
    )
    if explicit:
        return explicit
    return infer_decision_jobs(
        source_keys=keys,
        purpose_jobs=row.purpose_jobs,
        url=row.url,
        source_role=row.source_role,
    )


def infer_decision_jobs(
    *,
    source_keys: Sequence[str],
    purpose_jobs: str | None,
    url: str,
    source_role: str | None,
) -> list[str]:
    """Map inventory text into J01–J14. Conservative: unknown → empty (not J03)."""

    blob = " ".join(
        [
            " ".join(source_keys),
            purpose_jobs or "",
            url or "",
            source_role or "",
        ]
    ).lower()
    jobs: list[DecisionJob] = []

    def add(job: DecisionJob) -> None:
        if job not in jobs:
            jobs.append(job)

    if any(
        token in blob
        for token in (
            "universe",
            "isin",
            "instrument",
            "holiday",
            "calendar",
            "contract master",
            "equity_l",
        )
    ):
        add(DecisionJob.J01)
    if any(
        token in blob
        for token in (
            "mwpl",
            "ban",
            "asm",
            "gsm",
            "surveillance",
            "halt",
            "t2t",
            "price band",
            "tradability",
            "eligibility",
        )
    ):
        add(DecisionJob.J02)
    if any(
        token in blob
        for token in (
            "bhav",
            "quote",
            "price",
            "volume",
            "delivery",
            "depth",
            "liquidity",
            "ohlc",
            "candles",
        )
    ):
        add(DecisionJob.J03)
    if any(token in blob for token in ("preopen", "pre-open", "auction", "iep")):
        add(DecisionJob.J04)
    if any(
        token in blob
        for token in ("index", "breadth", "sector", "nifty", "banknifty", "regime")
    ):
        add(DecisionJob.J05)
    if any(
        token in blob
        for token in (
            "participant_oi",
            "participant-oi",
            "futures",
            "fii-dii",
            "basis",
            "rollover",
            "fo_bhav",
        )
    ):
        add(DecisionJob.J06)
    if any(
        token in blob
        for token in (
            "option",
            "pcr",
            "greeks",
            "max pain",
            "maxpain",
            "gex",
            "iv ",
            "implied",
        )
    ):
        add(DecisionJob.J07)
    if any(
        token in blob
        for token in ("insider", "sast", "pledge", "bulk", "block", "deal", "pit")
    ):
        add(DecisionJob.J08)
    if any(
        token in blob
        for token in (
            "corporate",
            "announcement",
            "filing",
            "result",
            "buyback",
            "takeover",
            "xbrl",
            "ca_",
        )
    ):
        add(DecisionJob.J09)
    if any(token in blob for token in ("amfi", "nsdl", "mutual fund", "portfolio")):
        add(DecisionJob.J10)
    if "mcx" in blob and any(
        token in blob
        for token in ("bhav", "future", "option", "delivery", "margin", "market watch")
    ):
        add(DecisionJob.J11)
    if any(
        token in blob
        for token in (
            "cftc",
            "lme",
            "shfe",
            "wgc",
            "eia",
            "opec",
            "baker",
            "usda",
            "fbil",
            "fred",
        )
    ):
        add(DecisionJob.J12)
    if any(
        token in blob for token in ("risk", "margin", "fee", "stt", "cost", "sizing")
    ):
        add(DecisionJob.J13)
    if any(
        token in blob
        for token in ("outcome", "calibration", "drift", "backtest", "label")
    ):
        add(DecisionJob.J14)

    if not jobs and source_role and "LOCAL" in source_role.upper():
        return []
    if not jobs and source_role and "REFERENCE" in source_role.upper():
        return []
    return [job.value for job in jobs]


def detect_row_defects(row: _RawInventoryRow) -> list[InventoryDefect]:
    defects: list[InventoryDefect] = []
    role = (row.source_role or "").upper()
    if any(
        token in role
        for token in ("SUPPORT", "REFERENCE", "OPEN_SOURCE", "LOCAL_APPLICATION")
    ):
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.NON_DATA_REFERENCE_ROW,
                inventoryId=row.inventory_id,
                url=row.url,
                detail="Non-data support, reference or local row requires quarantine",
            )
        )
    url = (row.url or "").strip()
    lowered = url.lower()

    if not url:
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.NON_HTTP_CANONICAL,
                inventoryId=row.inventory_id,
                url=url,
                detail="Empty canonical URL",
            )
        )
        return defects

    is_http = lowered.startswith(("http://", "https://"))
    key_prefixed = bool(re.match(r"^[a-z0-9_]+:https?://", url, re.I))
    is_sentinel = lowered in KNOWN_SENTINEL_URLS or (
        not is_http and ":no_snapshot" in lowered
    )

    if is_sentinel:
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.SENTINEL_URL,
                inventoryId=row.inventory_id,
                url=url,
                detail="Sentinel occupies canonical URL field",
            )
        )
    elif key_prefixed:
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.KEY_PREFIXED_URL,
                inventoryId=row.inventory_id,
                url=url,
                detail="URL cell carries source-key prefix before scheme",
            )
        )
    elif not is_http:
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.NON_HTTP_CANONICAL,
                inventoryId=row.inventory_id,
                url=url,
                detail="Canonical field is not an HTTP(S) URL",
            )
        )

    if "|" in url and ("http://" in lowered or "https://" in lowered):
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.COMPOUND_URL,
                inventoryId=row.inventory_id,
                url=url,
                detail="Compound URL cell joined by |",
            )
        )

    jobs = (row.purpose_jobs or "").strip()
    if not jobs or jobs.lower() in {"none", "null", "nan"}:
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.MISSING_PURPOSE_JOBS,
                inventoryId=row.inventory_id,
                url=url,
                detail="purpose_jobs missing",
            )
        )
    elif (
        len(jobs) > 80
        and " " in jobs
        and "," not in jobs
        and "|" not in jobs
        and not re.fullmatch(r"[A-Z0-9_\-\s/]+", jobs, flags=re.IGNORECASE)
    ):
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.PROSE_PURPOSE_JOBS,
                inventoryId=row.inventory_id,
                url=url,
                detail="purpose_jobs looks like free-form prose",
            )
        )

    keys_raw = row.active_source_keys
    if keys_raw and JSON_ARRAY_KEYS.match(str(keys_raw).strip()):
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.JSON_ARRAY_ACTIVE_KEYS,
                inventoryId=row.inventory_id,
                url=url,
                detail="active_source_keys uses JSON-array text",
            )
        )

    placeholder = (row.placeholder_or_test or "").strip().lower()
    role = (row.source_role or "").upper()
    if (
        placeholder in {"1", "true", "yes", "y"}
        or "TEST" in role
        or "PLACEHOLDER" in role
    ):
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.PLACEHOLDER_OR_TEST,
                inventoryId=row.inventory_id,
                url=url,
                detail="Placeholder/test row must be quarantined from activation metrics",
            )
        )

    if not (row.source_role or "").strip():
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.MISSING_ROLE,
                inventoryId=row.inventory_id,
                url=url,
                detail="source_role missing",
            )
        )

    return defects


def map_maturity_state(row: _RawInventoryRow) -> SourceMaturityState:
    """Advance only through explicit, contiguous proof steps.

    Free-form status prose is historical evidence, not an authorization contract.
    """

    proof = {
        item.upper() for item in _list_field(row.extra.get("maturity_proof_steps"))
    }
    current = SourceMaturityState.REGISTERED
    for state in MATURITY_LADDER[1:]:
        if state is SourceMaturityState.EXECUTION_AUTHORIZED:
            break
        if state.value not in proof:
            break
        current = state
    return current


def _publisher_from_url(url: str) -> str:
    host = (urlparse(url).hostname or "").lower()
    if not host:
        return "UNKNOWN"
    if host.startswith("www."):
        host = host[4:]
    return host


def _transport_variant(url: str) -> str:
    path = (urlparse(url).path or "").lower()
    if path.endswith((".csv", ".txt")):
        return "csv_download"
    if path.endswith((".zip",)):
        return "zip_download"
    if path.endswith((".xls", ".xlsx")):
        return "excel_download"
    if path.endswith((".pdf",)):
        return "pdf_document"
    if path.endswith((".xml",)):
        return "xml_feed"
    if "/api/" in path or "api." in (urlparse(url).hostname or ""):
        return "http_api"
    return "html_or_landing"


def _endpoint_role(url: str, source_role: str | None) -> str:
    role = (source_role or "").upper()
    if "SUPPORT" in role:
        return "support_referer"
    if "REFERENCE" in role:
        return "reference"
    if "LOCAL" in role:
        return "local_application"
    variant = _transport_variant(url)
    if variant in {"csv_download", "zip_download", "excel_download", "http_api"}:
        return "data_endpoint"
    return "landing_or_discovery"


def _evidence_family(jobs: Sequence[str], source_keys: Sequence[str]) -> str:
    keys = " ".join(source_keys).lower()
    if any(j in jobs for j in ("J07",)):
        return "OPTIONS_CHAIN"
    if any(j in jobs for j in ("J06",)) or "participant" in keys:
        return "DERIVATIVES_POSITIONING"
    if any(j in jobs for j in ("J02",)):
        return "TRADABILITY_SAFETY"
    if any(j in jobs for j in ("J08", "J09")):
        return "EVENTS_OWNERSHIP"
    if any(j in jobs for j in ("J11", "J12")):
        return "COMMODITY_CONTEXT"
    if any(j in jobs for j in ("J03", "J04")):
        return "PRICE_AND_LIQUIDITY"
    if any(j in jobs for j in ("J01",)):
        return "IDENTITY_UNIVERSE"
    if any(j in jobs for j in ("J10",)):
        return "DELAYED_INSTITUTIONAL"
    return "UNCLASSIFIED_OR_REFERENCE"


def build_dataset_root(
    *,
    root_id: str,
    urls: Sequence[str],
    source_keys: Sequence[str],
    source_role: str | None,
    jobs: Sequence[str],
    maturity: SourceMaturityState,
    linked: bool,
    overlap_resolutions: Sequence[OverlapResolution] = (),
) -> DatasetRootRecord:
    sample = list(dict.fromkeys(urls))[:3]
    primary_url = sample[0] if sample else ""
    mirror_group_id, resolver_id, fallback_cap = resolve_dataset_root_identity(
        urls=urls,
        source_keys=source_keys,
        overlap_resolutions=overlap_resolutions,
    )
    return DatasetRootRecord(
        datasetRootId=root_id,
        transportVariantId=_transport_variant(primary_url)
        if primary_url
        else "unknown",
        mirrorGroupId=mirror_group_id,
        resolverId=resolver_id,
        publisherAuthority=_publisher_from_url(primary_url)
        if primary_url
        else "UNKNOWN",
        endpointRole=_endpoint_role(primary_url, source_role),
        activationState=maturity.value,
        decisionJobs=list(jobs),
        evidenceFamily=_evidence_family(jobs, source_keys),
        familyWeightCap=0.0,
        fallbackAuthorityCap=fallback_cap
        if fallback_cap != "UNSPECIFIED_NO_VOTE"
        else "UNOFFICIAL_CANNOT_INHERIT_OFFICIAL",
        sourceKeys=list(dict.fromkeys(source_keys)),
        sampleUrls=sample,
        linked=linked,
        canUnlockExecution=False,
    )


def ban_mwpl_contract_splits() -> list[ContractSplitRecord]:
    """Hybrid §19.4 — split F&O ban hard veto from MWPL percentages."""

    return [
        ContractSplitRecord(
            contractId="nse_fno_ban",
            role="official_hard_veto",
            decisionJobs=[DecisionJob.J02.value],
            validEmptyAllowed=True,
            failureWaitReason="WAIT_SOURCE_SNAPSHOT",
            notes=(
                "Valid empty dated session means no listed bans. "
                "Fetch/stale/undated/schema failure must never mean no bans."
            ),
            runtimeAlias="nse_mwpl_ban",
        ),
        ContractSplitRecord(
            contractId="nse_mwpl_percentages",
            role="symbol_level_utilization_context",
            decisionJobs=[DecisionJob.J02.value, DecisionJob.J06.value],
            validEmptyAllowed=False,
            failureWaitReason="WAIT_MWPL_PERCENTAGES",
            notes=(
                "Not fully proven as a separate official percentage contract. "
                "Do not invent endpoints without live verification."
            ),
            runtimeAlias=None,
        ),
        ContractSplitRecord(
            contractId="fbil_usdinr_reference",
            role="daily_fx_context",
            decisionJobs=[DecisionJob.J12.value],
            validEmptyAllowed=False,
            failureWaitReason="WAIT_FX_CONTEXT",
            notes="Daily FBIL FX cannot unlock MCX intraday research confirm alone.",
            runtimeAlias=None,
        ),
        ContractSplitRecord(
            contractId="usd_inr_live",
            role="live_fx_via_licensed_or_openalgo_ro",
            decisionJobs=[DecisionJob.J11.value, DecisionJob.J12.value],
            validEmptyAllowed=False,
            failureWaitReason="WAIT_LIVE_FX",
            notes="Live FX only through approved read-only OpenAlgo/licensed path.",
            runtimeAlias=None,
        ),
    ]


def load_master_inventory_rows(
    workbook_path: Path | None = None,
) -> tuple[list[_RawInventoryRow], Path | None, str | None]:
    path = workbook_path or DEFAULT_WORKBOOK
    if not path.exists():
        return [], None, None

    from openpyxl import load_workbook  # type: ignore[import-untyped]

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "MASTER_CURRENT" not in wb.sheetnames:
            raise ValueError("MASTER_CURRENT sheet missing from inventory workbook")
        ws = wb["MASTER_CURRENT"]
        headers: list[str] | None = None
        rows: list[_RawInventoryRow] = []
        for index, values in enumerate(ws.iter_rows(values_only=True)):
            if index == 0:
                headers = [str(cell) if cell is not None else "" for cell in values]
                continue
            assert headers is not None
            data = {
                headers[i]: values[i] if i < len(values) else None
                for i in range(len(headers))
            }
            url = data.get("url") or data.get("master_canonical_url") or ""
            rows.append(
                _RawInventoryRow(
                    inventory_id=str(data.get("inventory_id") or index),
                    url=str(url).strip() if url is not None else "",
                    source_role=_as_optional_str(data.get("source_role")),
                    purpose_jobs=_as_optional_str(data.get("purpose_jobs")),
                    active_source_keys=_as_optional_str(data.get("active_source_keys")),
                    connection_status=_as_optional_str(data.get("connection_status")),
                    parser_status=_as_optional_str(data.get("parser_status")),
                    freshness_status=_as_optional_str(data.get("freshness_status")),
                    current_usability=_as_optional_str(data.get("current_usability")),
                    placeholder_or_test=_as_optional_str(
                        data.get("master_placeholder_flag")
                        or data.get("placeholder_or_test")
                    ),
                    sheet="MASTER_CURRENT",
                    extra=data,
                )
            )
    finally:
        wb.close()

    return rows, path, file_sha256(path)


def _as_optional_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text else None


def rows_from_mappings(items: Iterable[Mapping[str, Any]]) -> list[_RawInventoryRow]:
    rows: list[_RawInventoryRow] = []
    for index, item in enumerate(items, start=1):
        rows.append(
            _RawInventoryRow(
                inventory_id=str(item.get("inventory_id") or index),
                url=str(item.get("url") or "").strip(),
                source_role=_as_optional_str(item.get("source_role")),
                purpose_jobs=_as_optional_str(item.get("purpose_jobs")),
                active_source_keys=_as_optional_str(item.get("active_source_keys")),
                connection_status=_as_optional_str(item.get("connection_status")),
                parser_status=_as_optional_str(item.get("parser_status")),
                freshness_status=_as_optional_str(item.get("freshness_status")),
                current_usability=_as_optional_str(item.get("current_usability")),
                placeholder_or_test=_as_optional_str(item.get("placeholder_or_test")),
                sheet=str(item.get("sheet") or "FIXTURE"),
                extra=dict(item),
            )
        )
    return rows


def _extra_text(row: _RawInventoryRow, *names: str) -> str | None:
    for name in names:
        value = _as_optional_str(row.extra.get(name))
        if value:
            return value
    return None


def _canonical_input_value(row: _RawInventoryRow) -> str:
    return _extra_text(row, "master_canonical_url") or row.url


def _row_quarantine_reasons(
    row: _RawInventoryRow, urls: Sequence[str], jobs: Sequence[str]
) -> list[str]:
    reasons: list[str] = []
    role = (row.source_role or "").upper()
    if any(
        token in role
        for token in ("SUPPORT", "REFERENCE", "OPEN_SOURCE", "LOCAL_APPLICATION")
    ):
        reasons.append("NON_DATA_REFERENCE_ROW")
    placeholder = (row.placeholder_or_test or "").strip().lower()
    role = (row.source_role or "").upper()
    if placeholder in {"1", "true", "yes", "y"} or any(
        token in role for token in ("TEST", "PLACEHOLDER")
    ):
        reasons.append("PLACEHOLDER_OR_TEST")
    if not urls:
        reasons.append("NO_CANONICAL_ENDPOINT")
    if not row.source_role:
        reasons.append("MISSING_ROLE")
    if not jobs:
        reasons.append("MISSING_CONTROLLED_DECISION_JOB")
    return reasons


def _migration_actions_for_row(
    row: _RawInventoryRow,
    urls: Sequence[str],
    contract_ids: Sequence[str],
    defects: Sequence[InventoryDefect],
) -> list[MigrationAction]:
    actions: list[MigrationAction] = []
    action_by_defect = {
        InventoryDefectCode.NON_DATA_REFERENCE_ROW: "QUARANTINE_NON_DATA_REFERENCE",
        InventoryDefectCode.KEY_PREFIXED_URL: "NORMALIZE_KEY_PREFIX",
        InventoryDefectCode.COMPOUND_URL: "SPLIT_COMPOUND_URL",
        InventoryDefectCode.SENTINEL_URL: "MOVE_SENTINEL_TO_STATUS",
        InventoryDefectCode.NON_HTTP_CANONICAL: "QUARANTINE_NON_HTTP_VALUE",
        InventoryDefectCode.JSON_ARRAY_ACTIVE_KEYS: "NORMALIZE_SOURCE_KEYS",
        InventoryDefectCode.PLACEHOLDER_OR_TEST: "QUARANTINE_PLACEHOLDER",
        InventoryDefectCode.MISSING_ROLE: "REQUIRE_SOURCE_ROLE",
        InventoryDefectCode.MISSING_PURPOSE_JOBS: "REQUIRE_DECISION_JOB",
        InventoryDefectCode.PROSE_PURPOSE_JOBS: "NORMALIZE_DECISION_JOB",
    }
    seen: set[str] = set()
    for defect in defects:
        action_type = action_by_defect.get(defect.code)
        if not action_type or action_type in seen:
            continue
        seen.add(action_type)
        actions.append(
            MigrationAction(
                actionId=_stable_id(
                    "migration", f"{row.inventory_id}:{action_type}:{row.url}"
                ),
                inventoryId=row.inventory_id,
                actionType=action_type,
                originalValue=row.url,
                normalizedUrls=list(urls),
                sourceContractIds=list(contract_ids),
                reason=defect.detail,
                destructive=False,
            )
        )
    return actions


def _merged_extra(rows: Sequence[_RawInventoryRow], name: str) -> str | None:
    values = sorted(
        {
            value
            for row in rows
            if (value := _as_optional_str(row.extra.get(name))) is not None
        }
    )
    return values[0] if len(values) == 1 else None


def _reviewed_r3_binding(root_id: str) -> dict[str, Any] | None:
    """Return the approved R3 research binding; every other root fails closed."""
    if root_id != "nse_bhavcopy_eod":
        return None
    return {
        "feature_ids": ["FTR-040"],
        "independence_family": "PARTICIPATION",
        "authority_cap": "RESEARCH_DIRECTIONAL_WAIT_ONLY",
        "allowed_timeframes": ["SWING"],
        "allowed_instruments": ["NSE_EQ"],
        "directional_permission": True,
        "tests": ["test_r3_live.py"],
    }


def _compiled_source_contract(
    root_id: str, meta: Mapping[str, Any]
) -> CompiledSourceContract:
    rows = sorted(meta["rows"], key=lambda row: row.inventory_id)
    urls = sorted(meta["urls"])
    jobs = sorted(meta["jobs"])
    endpoint_ids = [_stable_id("endpoint", url) for url in urls]
    # TDG-GAP-001: the 15 AMEND-A-001 fields come only from reviewed
    # source-specific proofs. Workbook extras and catalog descriptors
    # are not publication calendars or schema versions.
    extended: dict[str, str] = {}
    from .source_extended_field_proofs import proof_for_source

    reviewed = proof_for_source(root_id)
    if reviewed is not None:
        for field, value in reviewed.fields.items():
            if field in EXTENDED_SOURCE_CONTRACT_FIELDS and value.strip():
                extended[field] = value.strip()
    missing = [
        field for field in EXTENDED_SOURCE_CONTRACT_FIELDS if field not in extended
    ]
    maturity = max(
        (map_maturity_state(row) for row in rows),
        key=MATURITY_LADDER.index,
        default=SourceMaturityState.REGISTERED,
    )
    quarantine = sorted(meta["quarantine"])
    safe_use = _merged_extra(rows, "safe_use") or (
        "Research discovery only; inspect source evidence before use"
    )
    blocked_use = (
        _merged_extra(rows, "blocked_use")
        or _merged_extra(rows, "why_not_currently_usable")
        or "Cannot unlock CONFIRMED, direction, quantity, execution or live-trade action"
    )
    next_action = _merged_extra(rows, "next_action") or (
        "Complete the source contract and fixture-backed parser validation"
    )
    reviewed_r3 = _reviewed_r3_binding(root_id)
    feature_ids = (
        list(reviewed_r3["feature_ids"])
        if reviewed_r3 is not None
        else _list_field(_merged_extra(rows, "feature_ids"))
    )
    complete = bool(urls and jobs and not missing and not quarantine)
    return CompiledSourceContract(
        sourceContractId=root_id,
        datasetRootId=root_id,
        endpointIds=endpoint_ids,
        inventoryIds=sorted({row.inventory_id for row in rows}),
        artifactKind=_merged_extra(rows, "artifact_kind")
        or (_transport_variant(urls[0]) if urls else "UNSPECIFIED"),
        parserId=_merged_extra(rows, "parser_id"),
        parserVersion=_merged_extra(rows, "parser_version"),
        datasetId=_merged_extra(rows, "dataset_id"),
        decisionJobs=jobs,
        featureIds=feature_ids,
        strategyProfiles=_list_field(_merged_extra(rows, "strategy_profiles")),
        mandatoryFor=_list_field(_merged_extra(rows, "mandatory_for")),
        confirmationFor=_list_field(_merged_extra(rows, "confirmation_for")),
        vetoFor=_list_field(_merged_extra(rows, "veto_for")),
        independenceFamily=(
            str(reviewed_r3["independence_family"])
            if reviewed_r3 is not None
            else _evidence_family(jobs, [root_id])
        ),
        authorityCap=(
            str(reviewed_r3["authority_cap"])
            if reviewed_r3 is not None
            else _merged_extra(rows, "authority_cap") or "UNSPECIFIED_NO_VOTE"
        ),
        freshnessByJob={job: "UNSPECIFIED" for job in jobs},
        fallbackChain=_list_field(_merged_extra(rows, "fallback_chain")),
        primaryPanelField=_merged_extra(rows, "primary_panel_field"),
        inspectorSections=_list_field(_merged_extra(rows, "inspector_sections")),
        tests=(
            list(reviewed_r3["tests"])
            if reviewed_r3 is not None
            else _list_field(_merged_extra(rows, "tests"))
        ),
        allowedTimeframes=(
            list(reviewed_r3["allowed_timeframes"])
            if reviewed_r3 is not None
            else _list_field(_merged_extra(rows, "allowed_timeframes"))
        ),
        allowedInstruments=(
            list(reviewed_r3["allowed_instruments"])
            if reviewed_r3 is not None
            else _list_field(_merged_extra(rows, "allowed_instruments"))
        ),
        directionalPermission=bool(
            reviewed_r3 is not None and reviewed_r3["directional_permission"]
        ),
        vetoPermission=False,
        quantityPermission=False,
        gatePermission=False,
        maturityState=maturity,
        maturityMissingPrerequisites=[
            state.value
            for state in MATURITY_LADDER[1:]
            if MATURITY_LADDER.index(state) > MATURITY_LADDER.index(maturity)
        ],
        extendedFields=extended,
        missingExtendedFields=missing,
        contractStatus=(
            "QUARANTINED" if quarantine else "COMPLETE" if complete else "INCOMPLETE"
        ),
        safeUse=safe_use,
        blockedUse=blocked_use,
        nextAction=next_action,
        linked=bool(meta["linked"]),
    )


def compile_inventory(
    rows: Sequence[_RawInventoryRow],
    *,
    workbook_path: Path | None = None,
    workbook_sha256: str | None = None,
    overlap_resolutions: Sequence[OverlapResolution] = (),
    expected_resolution_workbook_sha256: str | None = None,
    max_defects: int = 500,
    max_roots: int | None = None,
    compiler_state: CompilerState | None = None,
) -> InventoryCompilerReport:
    """Compile raw inventory rows into a lossless, fail-closed read-only view."""

    state = compiler_state or (CompilerState.OK if rows else CompilerState.INPUT_EMPTY)
    ok = state is CompilerState.OK and bool(rows)
    defects: list[InventoryDefect] = []
    migrations: list[MigrationAction] = []
    role_counts: Counter[str] = Counter()
    maturity_counts: Counter[str] = Counter(
        {maturity.value: 0 for maturity in MATURITY_LADDER}
    )
    roots: dict[str, dict[str, Any]] = {}
    endpoints: dict[str, dict[str, Any]] = {}
    linked_keys: set[str] = set()
    not_linked = 0
    sentinel_count = 0
    overlap_records: list[OverlapResolutionRecord] = []
    overlap_disposition_counts: Counter[str] = Counter(
        {disposition.value: 0 for disposition in OverlapDisposition}
    )
    resolution_by_url: dict[str, list[OverlapResolution]] = {}
    exact_resolution_keys: set[tuple[str, tuple[str, ...]]] = set()
    invalid_resolution_count = 0
    for resolution in overlap_resolutions:
        if resolution.exact_key in exact_resolution_keys:
            invalid_resolution_count += 1
        exact_resolution_keys.add(resolution.exact_key)
        resolution_by_url.setdefault(resolution.canonical_url, []).append(resolution)
    invalid_resolution_count += sum(
        max(0, len(items) - 1) for items in resolution_by_url.values()
    )
    workbook_resolution_mismatch = bool(
        expected_resolution_workbook_sha256
        and workbook_sha256 != expected_resolution_workbook_sha256
    )
    matched_resolution_keys: set[tuple[str, tuple[str, ...]]] = set()

    for row in rows:
        role_counts[row.source_role or "MISSING_ROLE"] += 1
        row_defects = detect_row_defects(row)
        defects.extend(row_defects)
        raw_urls = normalize_inventory_urls(_canonical_input_value(row))
        keys = sorted(set(parse_active_source_keys(row.active_source_keys)))
        contract_bindings = _source_contract_bindings(row, raw_urls)
        contract_ids = _source_contract_ids(row, raw_urls)
        jobs = _explicit_or_inferred_jobs(row, keys or contract_ids)
        quarantine = _row_quarantine_reasons(row, raw_urls, jobs)
        active_urls = raw_urls
        maturity = map_maturity_state(row)
        maturity_counts[maturity.value] += 1
        if not raw_urls and ":no_snapshot" in row.url.lower():
            sentinel_count += 1
        if keys:
            linked_keys.update(keys)
        else:
            not_linked += 1

        migrations.extend(
            _migration_actions_for_row(row, raw_urls, contract_ids, row_defects)
        )
        for contract_id in contract_ids:
            bucket = roots.setdefault(
                contract_id,
                {
                    "urls": set(),
                    "jobs": set(),
                    "rows": [],
                    "quarantine": set(),
                    "role": row.source_role,
                    "linked": bool(keys),
                },
            )
            bucket["urls"].update(
                url
                for url, bound_ids in contract_bindings.items()
                if contract_id in bound_ids
            )
            bucket["jobs"].update(jobs)
            bucket["rows"].append(row)
            bucket["quarantine"].update(quarantine)
            bucket["linked"] = bool(bucket["linked"] or keys)

        for url in active_urls:
            bound_contract_ids = contract_bindings.get(url, [])
            endpoint = endpoints.setdefault(
                url,
                {
                    "contract_ids": set(),
                    "inventory_ids": set(),
                    "partitions": set(),
                    "jobs": set(),
                    "rows": [],
                    "quarantine": set(),
                },
            )
            endpoint["contract_ids"].update(bound_contract_ids)
            endpoint["inventory_ids"].add(row.inventory_id)
            endpoint["partitions"].add(row.sheet)
            endpoint["jobs"].update(jobs)
            endpoint["rows"].append(row)
            endpoint["quarantine"].update(quarantine)

    for url, meta in sorted(endpoints.items()):
        inventory_ids = sorted(meta["inventory_ids"])
        contract_ids = sorted(meta["contract_ids"])
        if len(inventory_ids) <= 1:
            continue
        row_list = "|".join(inventory_ids)
        detail = f"Exact canonical URL appears in inventory rows {row_list}"
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.EXACT_URL_DUPLICATE,
                inventoryId=inventory_ids[0],
                url=url,
                detail=detail,
            )
        )
        migrations.append(
            MigrationAction(
                actionId=_stable_id("migration", f"duplicate:{url}"),
                inventoryId=inventory_ids[0],
                actionType="MERGE_EXACT_URL_DUPLICATE",
                originalValue=url,
                normalizedUrls=[url],
                sourceContractIds=contract_ids,
                reason=detail,
                destructive=False,
            )
        )

    for url, meta in sorted(endpoints.items()):
        inventory_ids = sorted(meta["inventory_ids"])
        contract_ids = sorted(meta["contract_ids"])
        if len(contract_ids) <= 1:
            continue
        contract_key = tuple(contract_ids)
        candidates = resolution_by_url.get(url, [])
        exact = [item for item in candidates if item.contract_ids == contract_key]
        resolution = (
            exact[0]
            if len(exact) == 1
            and len(candidates) == 1
            and not workbook_resolution_mismatch
            else None
        )
        if resolution is not None:
            matched_resolution_keys.add(resolution.exact_key)
            overlap_disposition_counts[resolution.disposition.value] += 1
            overlap_records.append(
                OverlapResolutionRecord(
                    canonicalUrl=url,
                    contractIds=contract_ids,
                    disposition=resolution.disposition,
                    state=resolution.state,
                    canonicalDatasetRootId=resolution.canonical_dataset_root_id,
                    inventoryIds=inventory_ids,
                    decisionJobs=sorted(meta["jobs"]),
                    evidenceReason=resolution.evidence_reason,
                    evidenceBasis=[item.value for item in resolution.evidence_basis],
                    resolutionVersion=resolution.resolution_version,
                    reviewedAt=resolution.reviewed_at,
                    gateEligible=False,
                )
            )
            if resolution.disposition is OverlapDisposition.QUARANTINED_UNPROVEN:
                meta["quarantine"].add("QUARANTINED_SEMANTIC_OVERLAP")
                for contract_id in contract_ids:
                    roots[contract_id]["quarantine"].add("QUARANTINED_SEMANTIC_OVERLAP")
            migrations.append(
                MigrationAction(
                    actionId=_stable_id("migration", f"resolved-overlap:{url}"),
                    inventoryId=inventory_ids[0],
                    actionType=f"RESOLVE_OVERLAP_{resolution.disposition.value}",
                    originalValue=url,
                    normalizedUrls=[url],
                    sourceContractIds=contract_ids,
                    reason=resolution.evidence_reason,
                    destructive=False,
                )
            )
            continue
        contract_list = "|".join(contract_ids)
        detail = f"Endpoint maps to unexplained source contracts {contract_list}"
        if candidates:
            detail += "; reviewed resolution is stale or conflicting"
        if workbook_resolution_mismatch:
            detail += "; reviewed workbook hash changed"
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.SEMANTIC_SOURCE_OVERLAP,
                inventoryId=inventory_ids[0],
                url=url,
                detail=detail,
            )
        )
        migrations.append(
            MigrationAction(
                actionId=_stable_id("migration", f"overlap:{url}"),
                inventoryId=inventory_ids[0],
                actionType="REVIEW_SEMANTIC_SOURCE_OVERLAP",
                originalValue=url,
                normalizedUrls=[url],
                sourceContractIds=contract_ids,
                reason=detail,
                destructive=False,
            )
        )
        overlap_records.append(
            OverlapResolutionRecord(
                canonicalUrl=url,
                contractIds=contract_ids,
                state=OverlapResolutionState.UNEXPLAINED_OVERLAP,
                inventoryIds=inventory_ids,
                decisionJobs=sorted(meta["jobs"]),
                evidenceReason=detail,
                gateEligible=False,
            )
        )

    stale_resolution_count = len(exact_resolution_keys - matched_resolution_keys)
    if workbook_resolution_mismatch:
        stale_resolution_count = max(stale_resolution_count, len(exact_resolution_keys))
    resolution_registry_valid = bool(
        invalid_resolution_count == 0 and stale_resolution_count == 0
    )

    catalog_keys = {item.key for item in SOURCE_CATALOG}
    expected_runtime_contracts = {"nse_fno_ban", "nse_mwpl_percentages"}
    missing_runtime_contracts = sorted(expected_runtime_contracts - catalog_keys)
    legacy_runtime_contracts = sorted({"nse_mwpl_ban"} & catalog_keys)
    if missing_runtime_contracts or legacy_runtime_contracts:
        missing_list = "|".join(missing_runtime_contracts) or "none"
        legacy_list = "|".join(legacy_runtime_contracts) or "none"
        detail = f"missing={missing_list}; legacy={legacy_list}"
        defects.append(
            InventoryDefect(
                code=InventoryDefectCode.REGISTRY_RUNTIME_CONTRACT_MISMATCH,
                inventoryId="RUNTIME_CONTRACT_MAP",
                url="",
                detail=detail,
            )
        )
        migrations.append(
            MigrationAction(
                actionId=_stable_id("migration", f"runtime-map:{detail}"),
                inventoryId="RUNTIME_CONTRACT_MAP",
                actionType="REMAP_RUNTIME_CONTRACT",
                originalValue="nse_mwpl_ban",
                normalizedUrls=[],
                sourceContractIds=sorted(expected_runtime_contracts),
                reason=detail,
                destructive=False,
            )
        )

    normalized_endpoints = [
        NormalizedEndpointRecord(
            endpointId=_stable_id("endpoint", url),
            canonicalUrl=url,
            sourceContractIds=sorted(meta["contract_ids"]),
            inventoryIds=sorted(meta["inventory_ids"]),
            partitionMemberships=sorted(meta["partitions"]),
            endpointRole=_endpoint_role(url, meta["rows"][0].source_role),
            purpose=(
                " | ".join(sorted(meta["jobs"]))
                if meta["jobs"]
                else "QUARANTINED_INVENTORY_REFERENCE"
            ),
            transportVariantId=_transport_variant(url),
            publisherAuthority=_publisher_from_url(url),
            decisionJobs=sorted(meta["jobs"]),
            safeUse=_merged_extra(meta["rows"], "safe_use")
            or "Research evidence only; source contract remains authoritative",
            blockedUse=_merged_extra(meta["rows"], "blocked_use")
            or _merged_extra(meta["rows"], "why_not_currently_usable")
            or "No execution, quantity or unsupported CONFIRMED state",
            nextAction=_merged_extra(meta["rows"], "next_action")
            or "Complete fixture-backed source-contract validation",
            quarantined=bool(meta["quarantine"]),
            quarantineReasons=sorted(meta["quarantine"]),
        )
        for url, meta in sorted(endpoints.items())
    ]
    source_contracts = [
        _compiled_source_contract(root_id, meta)
        for root_id, meta in sorted(roots.items())
    ]
    source_map = build_source_key_map(
        source_contracts,
        normalized_endpoints,
        SOURCE_CATALOG,
    )

    dataset_roots = [
        build_dataset_root(
            root_id=contract.source_contract_id,
            urls=[
                endpoint.canonical_url
                for endpoint in normalized_endpoints
                if endpoint.endpoint_id in contract.endpoint_ids
            ],
            source_keys=[contract.source_contract_id],
            source_role=roots[contract.source_contract_id]["role"],
            jobs=contract.decision_jobs,
            maturity=contract.maturity_state,
            linked=contract.linked,
            overlap_resolutions=overlap_resolutions,
        )
        for contract in source_contracts
    ]
    roots_truncated = max_roots is not None and len(dataset_roots) > max_roots
    if max_roots is not None:
        dataset_roots = dataset_roots[:max_roots]

    acceptance = _normalized_h1a0_acceptance(
        ok=ok,
        endpoints=normalized_endpoints,
        contracts=source_contracts,
        defects=defects,
        migrations=migrations,
        rows=rows,
        resolution_registry_valid=resolution_registry_valid,
        invalid_resolution_count=invalid_resolution_count,
        stale_resolution_count=stale_resolution_count,
    )
    passed = sum(item.passed for item in acceptance)
    defect_counts = Counter({code.value: 0 for code in InventoryDefectCode})
    defect_counts.update(defect.code.value for defect in defects)
    unresolved = sum(
        contract.contract_status != "COMPLETE" for contract in source_contracts
    )
    activation_ready = bool(
        ok
        and source_contracts
        and unresolved == 0
        and all(contract.gate_permission for contract in source_contracts)
    )

    return InventoryCompilerReport(
        compiledAt=datetime.now(timezone.utc),
        ok=ok,
        compilerState=state,
        workbookPath=str(workbook_path) if workbook_path else None,
        workbookSha256=workbook_sha256,
        rowCount=len(rows),
        roleCounts=dict(sorted(role_counts.items())),
        linkedSourceKeyCount=len(linked_keys),
        notLinkedRowCount=not_linked,
        activeCatalogKeyCount=len({item.key for item in SOURCE_CATALOG}),
        sourceKeyMap=source_map.records,
        endpointSourceMap=source_map.endpoint_records,
        sourceMapReviewVersion="CROSS-002-2026-09-03-v3",
        sourceMapReviewValid=source_map.review_valid,
        sourceMapReviewErrors=source_map.review_errors,
        sourceGovernanceReady=source_map.governance_ready,
        reviewedSourceKeyCount=source_map.reviewed_source_count,
        unreviewedSourceKeyCount=source_map.unreviewed_source_count,
        activationReviewedSourceKeyCount=(source_map.activation_reviewed_count),
        gateAuthorizedSourceKeyCount=source_map.gate_authorized_count,
        legacyUnresolvedCompiledIdentityCount=unresolved,
        namedInventorySourceKeyCount=source_map.named_inventory_count,
        generatedEndpointLineageCount=source_map.generated_lineage_count,
        runtimeCatalogSourceKeyCount=source_map.runtime_catalog_count,
        sharedSourceKeyCount=source_map.shared_count,
        inventoryOnlySourceKeyCount=source_map.inventory_only_count,
        runtimeOnlySourceKeyCount=source_map.runtime_only_count,
        livingSourceKeyCount=source_map.living_map_count,
        unassignedEndpointLineageCount=(source_map.unassigned_endpoint_count),
        namedInventorySourceKeySha256=(source_map.named_inventory_sha256),
        runtimeCatalogSourceKeySha256=(source_map.runtime_catalog_sha256),
        livingSourceKeySha256=source_map.living_map_sha256,
        defectCounts=dict(sorted(defect_counts.items())),
        h1a0DefectCodes=[code.value for code in HYBRID_H1A0_DEFECT_CODES],
        defects=defects[:max_defects],
        overlapResolutionRecords=overlap_records,
        overlapDispositionCounts=dict(sorted(overlap_disposition_counts.items())),
        invalidOverlapResolutionCount=invalid_resolution_count,
        staleOverlapResolutionCount=stale_resolution_count,
        resolutionRegistryValid=resolution_registry_valid,
        maturityStageCounts=dict(maturity_counts),
        decisionJobs=decision_job_catalog(),
        datasetRoots=dataset_roots,
        normalizedEndpoints=normalized_endpoints,
        sourceContracts=source_contracts,
        migrationManifest=migrations,
        inputCanonicalCount=len(
            {value for row in rows if (value := _canonical_input_value(row).strip())}
        ),
        normalizedEndpointCount=len(normalized_endpoints),
        normalizedSourceContractCount=len(source_contracts),
        sentinelStatusCount=sentinel_count,
        migrationActionCount=len(migrations),
        lineageRowCount=len(rows),
        rootsTruncated=roots_truncated,
        sourceActivationReady=activation_ready,
        unresolvedContractCount=unresolved,
        contractSplits=ban_mwpl_contract_splits(),
        extendedSourceContractFields=list(EXTENDED_SOURCE_CONTRACT_FIELDS),
        datasetRootIdentityFields=list(DATASET_ROOT_IDENTITY_FIELDS),
        acceptance=acceptance,
        h1a0AcceptancePassed=passed,
        h1a0AcceptanceTotal=len(acceptance),
        researchOnly=True,
        executionAuthorizedCount=0,
        rule=(
            "The compiler is a read-only normalization and contract-audit view. "
            "Linked, HTTP 200, parsed or fresh labels never authorize CONFIRMED, "
            "quantity, execution or live-trade action. Hash-generated source IDs "
            "are unassigned endpoint lineage, not named source contracts; the "
            "living source-key map and activation dispositions are independently "
            "drift-reviewed. Source governance readiness records a complete "
            "fail-closed review; it is separate from sourceActivationReady and "
            "never grants gate permission."
        ),
    )


def _normalized_h1a0_acceptance(
    *,
    ok: bool,
    endpoints: Sequence[NormalizedEndpointRecord],
    contracts: Sequence[CompiledSourceContract],
    defects: Sequence[InventoryDefect],
    migrations: Sequence[MigrationAction],
    rows: Sequence[_RawInventoryRow],
    resolution_registry_valid: bool,
    invalid_resolution_count: int,
    stale_resolution_count: int,
) -> list[AcceptanceCheck]:
    malformed = [
        item.canonical_url
        for item in endpoints
        if not item.canonical_url.startswith(("http://", "https://"))
    ]
    compound = [item.canonical_url for item in endpoints if "|" in item.canonical_url]
    sentinels = [
        item.canonical_url
        for item in endpoints
        if ":no_snapshot" in item.canonical_url.lower()
    ]
    lineage_ids = {item.inventory_id for item in migrations}
    endpoint_lineage = {
        inventory_id
        for endpoint in endpoints
        for inventory_id in endpoint.inventory_ids
    }
    covered_rows = lineage_ids | endpoint_lineage
    safe_contracts = all(
        contract.safe_use
        and contract.blocked_use
        and contract.next_action
        and (contract.decision_jobs or contract.contract_status == "QUARANTINED")
        for contract in contracts
    )
    canonical_rows_complete = safe_contracts and all(
        endpoint.endpoint_role
        and endpoint.purpose
        and endpoint.safe_use
        and endpoint.blocked_use
        and endpoint.next_action
        for endpoint in endpoints
    )
    defect_counts = Counter(defect.code.value for defect in defects)
    semantic_overlap_count = defect_counts[
        InventoryDefectCode.SEMANTIC_SOURCE_OVERLAP.value
    ]
    runtime_mismatch_count = defect_counts[
        InventoryDefectCode.REGISTRY_RUNTIME_CONTRACT_MISMATCH.value
    ]
    split_ids = {item.contract_id for item in ban_mwpl_contract_splits()}
    split_ok = {"nse_fno_ban", "nse_mwpl_percentages"} <= split_ids
    checks = [
        (
            "H1A0-01",
            "Canonical endpoints contain no source-key prefixes",
            not malformed,
            f"malformed={len(malformed)}",
        ),
        (
            "H1A0-02",
            "Every compiled endpoint contains exactly one URL",
            not compound,
            f"compound={len(compound)}",
        ),
        (
            "H1A0-03",
            "Sentinels are status records, not canonical endpoints",
            not sentinels,
            f"sentinels={len(sentinels)}",
        ),
        (
            "H1A0-04",
            "Every input row retains endpoint or migration lineage",
            len(covered_rows) == len(rows),
            f"covered={len(covered_rows)}; rows={len(rows)}",
        ),
        (
            "H1A0-05",
            "Every source has controlled jobs or quarantine plus explicit use limits",
            safe_contracts,
            f"contracts={len(contracts)}",
        ),
        (
            "H1A0-06",
            "F&O ban and MWPL percentages remain separate contracts",
            split_ok,
            f"split={split_ok}",
        ),
    ]
    checks = [
        (
            "H1A0-01",
            "Zero malformed source-key cells in the normalized view",
            not malformed,
            (
                f"normalized_malformed={len(malformed)}; "
                f"raw_key_prefix={defect_counts[InventoryDefectCode.KEY_PREFIXED_URL.value]}; "
                f"raw_json_keys={defect_counts[InventoryDefectCode.JSON_ARRAY_ACTIVE_KEYS.value]}"
            ),
        ),
        (
            "H1A0-02",
            "Zero compound URL cells in the normalized view",
            not compound,
            (
                f"normalized_compound={len(compound)}; "
                f"raw_compound={defect_counts[InventoryDefectCode.COMPOUND_URL.value]}"
            ),
        ),
        (
            "H1A0-03",
            "Zero sentinel values in canonical URL fields",
            not sentinels,
            (
                f"normalized_sentinels={len(sentinels)}; "
                f"raw_sentinels={defect_counts[InventoryDefectCode.SENTINEL_URL.value]}"
            ),
        ),
        (
            "H1A0-04",
            "Zero unexplained source-contract overlap with complete lineage",
            semantic_overlap_count == 0
            and len(covered_rows) == len(rows)
            and resolution_registry_valid,
            (
                f"unexplained_overlap={semantic_overlap_count}; "
                f"invalid_resolution={invalid_resolution_count}; "
                f"stale_resolution={stale_resolution_count}; "
                f"covered={len(covered_rows)}; rows={len(rows)}"
            ),
        ),
        (
            "H1A0-05",
            "Every canonical row has role, purpose, safe use, blocked use and next action",
            canonical_rows_complete,
            f"endpoints={len(endpoints)}; contracts={len(contracts)}",
        ),
        (
            "H1A0-06",
            "F&O ban and MWPL percentages remain separate runtime contracts",
            split_ok and runtime_mismatch_count == 0,
            f"split={split_ok}; runtime_mismatch={runtime_mismatch_count}",
        ),
    ]
    return [
        AcceptanceCheck(
            id=check_id,
            description=description,
            passed=bool(ok and passed),
            evidence=evidence if ok else f"compiler_input_failure; {evidence}",
        )
        for check_id, description, passed, evidence in checks
    ]


def compile_default_inventory(
    workbook_path: Path | None = None,
) -> InventoryCompilerReport:
    rows, path, digest = load_master_inventory_rows(workbook_path)
    if not rows:
        state = (
            CompilerState.INPUT_MISSING if path is None else CompilerState.INPUT_EMPTY
        )
        return compile_inventory(
            rows,
            workbook_path=path,
            workbook_sha256=digest,
            overlap_resolutions=REVIEWED_OVERLAP_RESOLUTIONS,
            expected_resolution_workbook_sha256=REVIEWED_WORKBOOK_SHA256,
            compiler_state=state,
        )
    return compile_inventory(
        rows,
        workbook_path=path,
        workbook_sha256=digest,
        overlap_resolutions=REVIEWED_OVERLAP_RESOLUTIONS,
        expected_resolution_workbook_sha256=REVIEWED_WORKBOOK_SHA256,
        compiler_state=CompilerState.OK,
    )


def decision_job_catalog() -> list[dict[str, str]]:
    return [
        {
            "id": job.value,
            "purpose": DECISION_JOB_SPECS[job]["purpose"],
            "scope": DECISION_JOB_SPECS[job]["scope"],
        }
        for job in DecisionJob
    ]
