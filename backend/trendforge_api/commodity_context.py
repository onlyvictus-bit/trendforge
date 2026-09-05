from __future__ import annotations

import json
import math
import re
import sqlite3
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, Literal
from urllib.parse import parse_qs, urlparse
from uuid import uuid4

from openpyxl import load_workbook
from pydantic import BaseModel, ConfigDict, Field
from pydantic.alias_generators import to_camel

from .institutional_sources import (
    ENDPOINTS,
    ContractStatus,
    EndpointFetchResult,
    FetchState,
)


ROOT_DIR = Path(__file__).resolve().parents[2]
DEFAULT_DB = ROOT_DIR / "data" / "trendforge_research.db"
COMMODITY_ENDPOINTS = {
    "mcx_option_chain",
    "mcx_market_watch",
    "mcx_delivery_reports",
    "sge_benchmark_gold",
    "baker_hughes_na_rig_count",
}


class CommodityModel(BaseModel):
    model_config = ConfigDict(
        alias_generator=to_camel,
        populate_by_name=True,
        serialize_by_alias=True,
    )


class CommoditySourceStatus(CommodityModel):
    endpoint_key: str
    fetch_state: str
    parser_state: Literal[
        "STRUCTURED_OK",
        "STRUCTURED_RESEARCH_ONLY",
        "VALID_EMPTY",
        "STALE_FALLBACK",
        "SCHEMA_MISMATCH",
        "FETCH_FAILED",
    ]
    fetched_at: datetime
    content_hash: str | None = None
    record_count: int = Field(default=0, ge=0)
    normalized_count: int = Field(default=0, ge=0)
    contract_status: str
    reason: str


class MCXOptionRow(CommodityModel):
    source_content_hash: str | None = None
    data_date: str | None = None
    symbol: str
    expiry: str
    strike: float
    option_type: Literal["CE", "PE"]
    open_interest: float = Field(ge=0)
    change_in_oi: float
    volume: float = Field(ge=0)
    last_price: float = Field(ge=0)
    bid_quantity: float = Field(ge=0)
    bid_price: float = Field(ge=0)
    ask_quantity: float = Field(ge=0)
    ask_price: float = Field(ge=0)
    underlying_value: float | None = None


class MCXOptionMetrics(CommodityModel):
    symbol: str
    expiry: str
    data_date: str | None = None
    underlying_value: float | None = None
    strike_count: int = Field(ge=0)
    call_open_interest: float = Field(ge=0)
    put_open_interest: float = Field(ge=0)
    pcr_oi: float | None = None
    call_wall_strike: float | None = None
    put_wall_strike: float | None = None
    max_pain_strike: float | None = None
    total_volume: float = Field(ge=0)
    signal_state: Literal["CONTEXT_AVAILABLE", "WAIT_NO_OPEN_INTEREST"]
    limitations: list[str] = Field(default_factory=list)
    can_unlock_ready: bool = False


class MCXMarketWatchRow(CommodityModel):
    source_content_hash: str | None = None
    data_date: str | None = None
    symbol: str
    product_code: str | None = None
    expiry: str | None = None
    instrument_name: str | None = None
    open_price: float | None = None
    high_price: float | None = None
    low_price: float | None = None
    last_price: float = Field(ge=0)
    previous_close: float | None = None
    percent_change: float | None = None
    volume: float = Field(ge=0)
    open_interest: float = Field(ge=0)
    buy_price: float | None = None
    sell_price: float | None = None
    buy_quantity: float | None = None
    sell_quantity: float | None = None
    bid_ask_spread: float | None = None
    can_unlock_ready: bool = False


class SGEObservation(CommodityModel):
    observation_date: str
    series_key: Literal["zp", "wp"]
    value: float
    source_content_hash: str | None = None


class SGEBenchmarkMetrics(CommodityModel):
    data_date: str
    latest_zp: float | None = None
    latest_wp: float | None = None
    zp_change_5_observations_pct: float | None = None
    wp_change_5_observations_pct: float | None = None
    intraday_spread_proxy: float | None = None
    lbma_premium: float | None = None
    trend_state: Literal["RISING", "FALLING", "MIXED", "UNKNOWN"]
    limitation: str
    can_unlock_ready: bool = False


class BakerHughesMetrics(CommodityModel):
    data_date: str
    us_total: int
    canada_total: int
    north_america_total: int
    north_america_weekly_change: int
    north_america_yearly_change: int
    us_oil: int | None = None
    us_gas: int | None = None
    supply_signal: Literal[
        "SUPPLY_EXPANSION", "SUPPLY_CONTRACTION", "SUPPLY_STABLE"
    ]
    limitation: str = "Rig counts are delayed supply context, not an intraday crude trigger."
    can_unlock_ready: bool = False


class CommodityContextSnapshot(CommodityModel):
    run_id: str
    normalized_at: datetime
    state: Literal["RESEARCH_ONLY", "WAIT_SOURCE", "WAIT_NO_SNAPSHOT"]
    sources: list[CommoditySourceStatus] = Field(default_factory=list)
    mcx_option_rows: list[MCXOptionRow] = Field(default_factory=list)
    mcx_option_metrics: list[MCXOptionMetrics] = Field(default_factory=list)
    mcx_market_rows: list[MCXMarketWatchRow] = Field(default_factory=list)
    sge_observations: list[SGEObservation] = Field(default_factory=list)
    sge_benchmark: SGEBenchmarkMetrics | None = None
    baker_hughes: BakerHughesMetrics | None = None
    can_unlock_ready: bool = False
    required_confirmations: list[str] = Field(default_factory=list)
    reason: str


def _number(value: Any, default: float = 0.0) -> float:
    if value in (None, "", "-"):
        return default
    try:
        parsed = float(str(value).replace(",", ""))
    except (TypeError, ValueError):
        return default
    return parsed if math.isfinite(parsed) else default


def _dotnet_date(value: Any) -> str | None:
    match = re.search(r"Date\((-?\d+)", str(value or ""))
    if not match:
        return None
    try:
        return datetime.fromtimestamp(int(match.group(1)) / 1000, tz=UTC).date().isoformat()
    except (OverflowError, OSError, ValueError):
        return None


def _date_from_text(value: Any) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    for format_string in (
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%d-%b-%Y %H:%M:%S",
        "%d%b%Y",
        "%d-%b-%Y",
    ):
        try:
            return datetime.strptime(text, format_string).date().isoformat()
        except ValueError:
            continue
    return _dotnet_date(text)


def _payload_rows(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, list):
        return [row for row in payload if isinstance(row, dict)]
    if not isinstance(payload, dict):
        return []
    for key in ("Data", "data", "records", "rows", "Table"):
        value = payload.get(key)
        if isinstance(value, list):
            return [row for row in value if isinstance(row, dict)]
        if isinstance(value, dict):
            nested = _payload_rows(value)
            if nested:
                return nested
    return []


def _query_parameter(url: str, name: str) -> str:
    values = parse_qs(urlparse(url).query)
    target = name.lower()
    for key, value in values.items():
        if key.lower() == target and value:
            return str(value[0]).strip().upper()
    return ""


def _text(row: dict[str, Any], *keys: str) -> str:
    for key in keys:
        value = row.get(key)
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _max_pain(rows: list[MCXOptionRow], underlying: float | None) -> float | None:
    call_oi = {
        row.strike: row.open_interest
        for row in rows
        if row.option_type == "CE" and row.open_interest > 0
    }
    put_oi = {
        row.strike: row.open_interest
        for row in rows
        if row.option_type == "PE" and row.open_interest > 0
    }
    strikes = sorted(set(call_oi) | set(put_oi))
    if not strikes:
        return None

    def writer_payout(settlement: float) -> float:
        calls = sum(
            max(settlement - strike, 0) * open_interest
            for strike, open_interest in call_oi.items()
        )
        puts = sum(
            max(strike - settlement, 0) * open_interest
            for strike, open_interest in put_oi.items()
        )
        return calls + puts

    anchor = underlying if underlying is not None else strikes[len(strikes) // 2]
    return min(strikes, key=lambda strike: (writer_payout(strike), abs(strike - anchor)))


def _normalize_mcx(
    result: EndpointFetchResult,
) -> tuple[list[MCXOptionRow], MCXOptionMetrics | None, str | None]:
    payload = result.payload
    if not isinstance(payload, dict) or not isinstance(payload.get("Data"), list):
        return [], None, "MCX option-chain payload lacks a Data list."
    symbol = _query_parameter(result.url, "Symbol")
    expiry = _query_parameter(result.url, "Expiry")
    if not symbol or not expiry:
        return [], None, "MCX option-chain URL lacks Symbol or Expiry lineage."
    summary = payload.get("Summary") if isinstance(payload.get("Summary"), dict) else {}
    data_date = _dotnet_date(summary.get("AsOn"))
    rows: list[MCXOptionRow] = []
    strikes: set[float] = set()
    underlying_values: list[float] = []
    for item in payload["Data"]:
        if not isinstance(item, dict):
            continue
        strike = _number(
            item.get("CE_StrikePrice")
            or item.get("PE_StrikePrice")
            or item.get("StrikePrice"),
            default=-1,
        )
        if strike < 0:
            continue
        strikes.add(strike)
        underlying = _number(item.get("UnderlyingValue"), default=-1)
        if underlying >= 0:
            underlying_values.append(underlying)
        for option_type in ("CE", "PE"):
            rows.append(
                MCXOptionRow(
                    source_content_hash=result.content_hash,
                    data_date=data_date,
                    symbol=symbol,
                    expiry=expiry,
                    strike=strike,
                    option_type=option_type,
                    open_interest=max(_number(item.get(f"{option_type}_OpenInterest")), 0),
                    change_in_oi=_number(item.get(f"{option_type}_ChangeInOI")),
                    volume=max(_number(item.get(f"{option_type}_Volume")), 0),
                    last_price=max(_number(item.get(f"{option_type}_LTP")), 0),
                    bid_quantity=max(_number(item.get(f"{option_type}_BidQty")), 0),
                    bid_price=max(_number(item.get(f"{option_type}_BidPrice")), 0),
                    ask_quantity=max(_number(item.get(f"{option_type}_AskQty")), 0),
                    ask_price=max(_number(item.get(f"{option_type}_AskPrice")), 0),
                    underlying_value=underlying if underlying >= 0 else None,
                )
            )
    if not rows:
        return [], None, "MCX option-chain payload contains no valid strike rows."
    call_rows = [row for row in rows if row.option_type == "CE"]
    put_rows = [row for row in rows if row.option_type == "PE"]
    call_oi = sum(row.open_interest for row in call_rows)
    put_oi = sum(row.open_interest for row in put_rows)
    underlying = underlying_values[-1] if underlying_values else None
    has_oi = call_oi > 0 or put_oi > 0
    metrics = MCXOptionMetrics(
        symbol=symbol,
        expiry=expiry,
        data_date=data_date,
        underlying_value=underlying,
        strike_count=len(strikes),
        call_open_interest=call_oi,
        put_open_interest=put_oi,
        pcr_oi=round(put_oi / call_oi, 6) if call_oi > 0 else None,
        call_wall_strike=(
            max(call_rows, key=lambda row: row.open_interest).strike
            if call_oi > 0
            else None
        ),
        put_wall_strike=(
            max(put_rows, key=lambda row: row.open_interest).strike
            if put_oi > 0
            else None
        ),
        max_pain_strike=_max_pain(rows, underlying) if has_oi else None,
        total_volume=sum(row.volume for row in rows),
        signal_state="CONTEXT_AVAILABLE" if has_oi else "WAIT_NO_OPEN_INTEREST",
        limitations=[
            "The public endpoint contract is unverified and remains research-only.",
            "PCR and max pain are positioning context, not directional trade triggers.",
            "Greeks and gamma exposure are unavailable unless IV, lot size, and signed positioning assumptions are validated.",
        ],
    )
    return rows, metrics, None


def _normalize_mcx_market_watch(
    result: EndpointFetchResult,
) -> tuple[list[MCXMarketWatchRow], str | None]:
    rows: list[MCXMarketWatchRow] = []
    for item in _payload_rows(result.payload):
        symbol = _text(item, "Symbol", "ProductCode", "Commodity")
        if not symbol:
            continue
        buy_price = _number(item.get("BuyPrice"), default=-1)
        sell_price = _number(item.get("SellPrice"), default=-1)
        rows.append(
            MCXMarketWatchRow(
                source_content_hash=result.content_hash,
                data_date=_date_from_text(
                    item.get("LTTValue")
                    or item.get("LastUpdatedTime")
                    or item.get("TradeDate")
                    or item.get("ExpiryDate")
                ),
                symbol=symbol.upper(),
                product_code=_text(item, "ProductCode") or None,
                expiry=_text(item, "ExpiryDate", "Expiry") or None,
                instrument_name=_text(item, "InstrumentName", "Instrument") or None,
                open_price=_number(item.get("Open"), default=-1) if item.get("Open") not in (None, "") else None,
                high_price=_number(item.get("High"), default=-1) if item.get("High") not in (None, "") else None,
                low_price=_number(item.get("Low"), default=-1) if item.get("Low") not in (None, "") else None,
                last_price=max(_number(item.get("LTP") or item.get("LastTradedPrice"), default=0), 0),
                previous_close=_number(item.get("PreviousClose"), default=-1) if item.get("PreviousClose") not in (None, "") else None,
                percent_change=_number(item.get("PercentChange") or item.get("AbsoluteChange"), default=-1) if (item.get("PercentChange") or item.get("AbsoluteChange")) not in (None, "") else None,
                volume=max(_number(item.get("Volume") or item.get("VolumeInLots"), default=0), 0),
                open_interest=max(_number(item.get("OpenInterest") or item.get("OI"), default=0), 0),
                buy_price=buy_price if buy_price >= 0 else None,
                sell_price=sell_price if sell_price >= 0 else None,
                buy_quantity=_number(item.get("BuyQuantity") or item.get("BidQty"), default=-1) if (item.get("BuyQuantity") or item.get("BidQty")) not in (None, "") else None,
                sell_quantity=_number(item.get("SellQuantity") or item.get("AskQty"), default=-1) if (item.get("SellQuantity") or item.get("AskQty")) not in (None, "") else None,
                bid_ask_spread=round(sell_price - buy_price, 6) if sell_price >= 0 and buy_price >= 0 else None,
            )
        )
    if not rows:
        return [], "MCX market-watch payload contains no valid contract rows."
    return rows, None


def _series(payload: dict[str, Any], key: str, content_hash: str | None) -> list[SGEObservation]:
    output: list[SGEObservation] = []
    raw = payload.get(key)
    if not isinstance(raw, list):
        return output
    for item in raw:
        if not isinstance(item, list) or len(item) < 2:
            continue
        try:
            observed = datetime.fromtimestamp(float(item[0]) / 1000, tz=UTC).date()
            value = float(item[1])
        except (TypeError, ValueError, OSError, OverflowError):
            continue
        if math.isfinite(value):
            output.append(
                SGEObservation(
                    observation_date=observed.isoformat(),
                    series_key=key,
                    value=value,
                    source_content_hash=content_hash,
                )
            )
    output.sort(key=lambda row: row.observation_date)
    return output


def _change(rows: list[SGEObservation], lookback: int = 5) -> float | None:
    if len(rows) <= lookback or rows[-lookback - 1].value == 0:
        return None
    return round((rows[-1].value / rows[-lookback - 1].value - 1) * 100, 4)


def _normalize_sge(
    result: EndpointFetchResult,
) -> tuple[list[SGEObservation], SGEBenchmarkMetrics | None, str | None]:
    if not isinstance(result.payload, dict):
        return [], None, "SGE benchmark payload is not an object."
    zp = _series(result.payload, "zp", result.content_hash)
    wp = _series(result.payload, "wp", result.content_hash)
    observations = zp + wp
    if not observations:
        return [], None, "SGE benchmark payload contains no valid zp/wp observations."
    zp_change = _change(zp)
    wp_change = _change(wp)
    changes = [value for value in (zp_change, wp_change) if value is not None]
    if changes and all(value > 0 for value in changes):
        trend = "RISING"
    elif changes and all(value < 0 for value in changes):
        trend = "FALLING"
    elif changes:
        trend = "MIXED"
    else:
        trend = "UNKNOWN"
    latest_zp = zp[-1].value if zp else None
    latest_wp = wp[-1].value if wp else None
    dates = [row.observation_date for row in observations]
    metrics = SGEBenchmarkMetrics(
        data_date=max(dates),
        latest_zp=latest_zp,
        latest_wp=latest_wp,
        zp_change_5_observations_pct=zp_change,
        wp_change_5_observations_pct=wp_change,
        intraday_spread_proxy=(
            round(latest_wp - latest_zp, 6)
            if latest_zp is not None and latest_wp is not None
            else None
        ),
        trend_state=trend,
        limitation=(
            "This source contains SGE benchmark series only. LBMA price, FX, tax, "
            "unit conversion, and timestamp alignment are absent, so no SGE/LBMA premium is calculated."
        ),
    )
    return observations, metrics, None


def _parse_workbook_date(value: Any) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    for format_string in ("%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(str(value).strip(), format_string).date().isoformat()
        except ValueError:
            continue
    return None


def _normalize_baker(
    result: EndpointFetchResult,
) -> tuple[BakerHughesMetrics | None, str | None]:
    path = Path(result.raw_path or "")
    if not path.is_file():
        return None, "Baker Hughes raw workbook path is missing."
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return None, f"Baker Hughes workbook could not be opened: {type(exc).__name__}."
    if "NAM Summary" not in workbook.sheetnames:
        workbook.close()
        return None, "Baker Hughes workbook lacks the NAM Summary sheet."
    sheet = workbook["NAM Summary"]
    data_date = _parse_workbook_date(sheet["D4"].value)
    rows: dict[str, tuple[int, int, int, int, int]] = {}
    section = "TOTALS"
    for values in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 100), values_only=True):
        label = str(values[1] or "").strip() if len(values) > 1 else ""
        if label == "U.S. Breakout Information":
            section = "US"
            continue
        if label == "Canada Breakout Information":
            section = "CANADA"
            continue
        if label in {"United States Total", "Canada", "North America"} or (
            section == "US" and label in {"Oil", "Gas"}
        ):
            rows[f"{section}:{label}"] = (
                int(_number(values[3] if len(values) > 3 else None)),
                int(_number(values[4] if len(values) > 4 else None)),
                int(_number(values[5] if len(values) > 5 else None)),
                int(_number(values[6] if len(values) > 6 else None)),
                int(_number(values[7] if len(values) > 7 else None)),
            )
    workbook.close()
    required = {
        "TOTALS:United States Total",
        "TOTALS:Canada",
        "TOTALS:North America",
    }
    if not data_date or not required <= rows.keys():
        return None, "Baker Hughes summary lacks date or required regional totals."
    us = rows["TOTALS:United States Total"]
    canada = rows["TOTALS:Canada"]
    north_america = rows["TOTALS:North America"]
    weekly_change = north_america[1]
    signal = (
        "SUPPLY_EXPANSION"
        if weekly_change > 0
        else "SUPPLY_CONTRACTION"
        if weekly_change < 0
        else "SUPPLY_STABLE"
    )
    return BakerHughesMetrics(
        data_date=data_date,
        us_total=us[0],
        canada_total=canada[0],
        north_america_total=north_america[0],
        north_america_weekly_change=weekly_change,
        north_america_yearly_change=north_america[3],
        us_oil=rows.get("US:Oil", (0, 0, 0, 0, 0))[0] or None,
        us_gas=rows.get("US:Gas", (0, 0, 0, 0, 0))[0] or None,
        supply_signal=signal,
    ), None


def _status(
    result: EndpointFetchResult,
    *,
    normalized_count: int,
    error: str | None,
) -> CommoditySourceStatus:
    contract = ENDPOINTS[result.endpoint_key].contract_status
    if result.state == FetchState.NO_DATA_NOW:
        parser_state = "VALID_EMPTY"
        reason = "Endpoint returned a valid empty response for this request."
    elif result.state not in {FetchState.RAW_ARCHIVED, FetchState.STALE_FALLBACK}:
        parser_state = "FETCH_FAILED"
        reason = result.reason or "Fetch did not produce a usable archived artifact."
    elif error:
        parser_state = "SCHEMA_MISMATCH"
        reason = error
    elif result.state == FetchState.STALE_FALLBACK:
        parser_state = "STALE_FALLBACK"
        reason = "A cached artifact was normalized but cannot confirm current conditions."
    elif contract == ContractStatus.UNVERIFIED_RESEARCH:
        parser_state = "STRUCTURED_RESEARCH_ONLY"
        reason = "Rows normalized with lineage; endpoint contract remains unverified and cannot unlock READY."
    else:
        parser_state = "STRUCTURED_OK"
        reason = "Official artifact normalized with source-specific schema and lineage."
    return CommoditySourceStatus(
        endpoint_key=result.endpoint_key,
        fetch_state=result.state.value,
        parser_state=parser_state,
        fetched_at=result.fetched_at,
        content_hash=result.content_hash,
        record_count=result.record_count,
        normalized_count=normalized_count,
        contract_status=contract.value,
        reason=reason,
    )


def build_commodity_context_snapshot(
    results: list[EndpointFetchResult],
) -> CommodityContextSnapshot:
    sources: list[CommoditySourceStatus] = []
    mcx_rows: list[MCXOptionRow] = []
    mcx_metrics: list[MCXOptionMetrics] = []
    mcx_market_rows: list[MCXMarketWatchRow] = []
    sge_rows: list[SGEObservation] = []
    sge_metrics: SGEBenchmarkMetrics | None = None
    baker_metrics: BakerHughesMetrics | None = None
    for result in results:
        if result.endpoint_key not in COMMODITY_ENDPOINTS:
            continue
        normalized_count = 0
        error: str | None = None
        if result.state in {FetchState.RAW_ARCHIVED, FetchState.STALE_FALLBACK}:
            if result.endpoint_key == "mcx_option_chain":
                rows, metrics, error = _normalize_mcx(result)
                mcx_rows.extend(rows)
                if metrics:
                    mcx_metrics.append(metrics)
                normalized_count = len(rows)
            elif result.endpoint_key == "mcx_market_watch":
                rows, error = _normalize_mcx_market_watch(result)
                mcx_market_rows.extend(rows)
                normalized_count = len(rows)
            elif result.endpoint_key == "mcx_delivery_reports":
                error = "MCX delivery report index archived; Excel link discovery/parser is pending."
            elif result.endpoint_key == "sge_benchmark_gold":
                rows, metrics, error = _normalize_sge(result)
                sge_rows.extend(rows)
                sge_metrics = metrics
                normalized_count = len(rows)
            elif result.endpoint_key == "baker_hughes_na_rig_count":
                baker_metrics, error = _normalize_baker(result)
                normalized_count = 1 if baker_metrics else 0
        sources.append(
            _status(result, normalized_count=normalized_count, error=error)
        )
    has_context = bool(mcx_metrics or sge_metrics or baker_metrics)
    return CommodityContextSnapshot(
        run_id=str(uuid4()),
        normalized_at=datetime.now(UTC),
        state="RESEARCH_ONLY" if has_context else "WAIT_SOURCE",
        sources=sources,
        mcx_option_rows=mcx_rows,
        mcx_option_metrics=mcx_metrics,
        mcx_market_rows=mcx_market_rows,
        sge_observations=sge_rows,
        sge_benchmark=sge_metrics,
        baker_hughes=baker_metrics,
        required_confirmations=[
            "Fresh official MCX bhavcopy price, volume and open interest",
            "Global benchmark direction and USD/INR translation",
            "Fresh contract-specific structure, liquidity, spread and risk checks",
            "CFTC positioning and commodity event context where applicable",
        ],
        reason=(
            "Commodity source data is normalized for research context. It cannot independently unlock READY."
            if has_context
            else "No commodity artifact passed source-specific normalization."
        ),
    )


def _connect(path: Path) -> sqlite3.Connection:
    path.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(path, timeout=10)
    connection.row_factory = sqlite3.Row
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA busy_timeout=5000")
    return connection


def _initialize(connection: sqlite3.Connection) -> None:
    connection.executescript(
        """
        CREATE TABLE IF NOT EXISTS commodity_context_runs (
            run_id TEXT PRIMARY KEY,
            normalized_at TEXT NOT NULL,
            state TEXT NOT NULL,
            source_count INTEGER NOT NULL,
            payload_json TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS mcx_option_chain_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id TEXT NOT NULL,
            source_content_hash TEXT,
            data_date TEXT,
            symbol TEXT NOT NULL,
            expiry TEXT NOT NULL,
            strike REAL NOT NULL,
            option_type TEXT NOT NULL,
            open_interest REAL NOT NULL,
            change_in_oi REAL NOT NULL,
            volume REAL NOT NULL,
            last_price REAL NOT NULL,
            bid_quantity REAL NOT NULL,
            bid_price REAL NOT NULL,
            ask_quantity REAL NOT NULL,
            ask_price REAL NOT NULL,
            underlying_value REAL,
            UNIQUE(run_id, symbol, expiry, strike, option_type)
        );
        CREATE TABLE IF NOT EXISTS sge_benchmark_observations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            run_id TEXT NOT NULL,
            source_content_hash TEXT,
            observation_date TEXT NOT NULL,
            series_key TEXT NOT NULL,
            value REAL NOT NULL,
            UNIQUE(run_id, observation_date, series_key)
        );
        CREATE INDEX IF NOT EXISTS idx_commodity_context_latest
        ON commodity_context_runs(normalized_at DESC);
        CREATE INDEX IF NOT EXISTS idx_mcx_option_context
        ON mcx_option_chain_rows(symbol, expiry, data_date, strike);
        """
    )


def save_commodity_context_snapshot(
    snapshot: CommodityContextSnapshot, db_path: Path = DEFAULT_DB
) -> None:
    payload = snapshot.model_dump(mode="json", by_alias=True)
    with _connect(db_path) as connection:
        _initialize(connection)
        connection.execute(
            """
            INSERT INTO commodity_context_runs(
                run_id, normalized_at, state, source_count, payload_json
            ) VALUES (?, ?, ?, ?, ?)
            """,
            (
                snapshot.run_id,
                snapshot.normalized_at.isoformat(),
                snapshot.state,
                len(snapshot.sources),
                json.dumps(payload, ensure_ascii=True, separators=(",", ":")),
            ),
        )
        connection.executemany(
            """
            INSERT INTO mcx_option_chain_rows(
                run_id, source_content_hash, data_date, symbol, expiry, strike,
                option_type, open_interest, change_in_oi, volume, last_price,
                bid_quantity, bid_price, ask_quantity, ask_price, underlying_value
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                (
                    snapshot.run_id,
                    row.source_content_hash,
                    row.data_date,
                    row.symbol,
                    row.expiry,
                    row.strike,
                    row.option_type,
                    row.open_interest,
                    row.change_in_oi,
                    row.volume,
                    row.last_price,
                    row.bid_quantity,
                    row.bid_price,
                    row.ask_quantity,
                    row.ask_price,
                    row.underlying_value,
                )
                for row in snapshot.mcx_option_rows
            ],
        )
        connection.executemany(
            """
            INSERT INTO sge_benchmark_observations(
                run_id, source_content_hash, observation_date, series_key, value
            ) VALUES (?, ?, ?, ?, ?)
            """,
            [
                (
                    snapshot.run_id,
                    row.source_content_hash,
                    row.observation_date,
                    row.series_key,
                    row.value,
                )
                for row in snapshot.sge_observations
            ],
        )


def latest_commodity_context_snapshot(
    *, db_path: Path = DEFAULT_DB
) -> CommodityContextSnapshot:
    with _connect(db_path) as connection:
        _initialize(connection)
        row = connection.execute(
            "SELECT payload_json FROM commodity_context_runs ORDER BY normalized_at DESC LIMIT 1"
        ).fetchone()
    if row is None:
        return CommodityContextSnapshot(
            run_id="NO_SNAPSHOT",
            normalized_at=datetime.now(UTC),
            state="WAIT_NO_SNAPSHOT",
            reason="Run the extended market-source refresh to create commodity context.",
        )
    return CommodityContextSnapshot.model_validate_json(row["payload_json"])


def normalize_and_save_commodity_context(
    results: list[EndpointFetchResult], db_path: Path = DEFAULT_DB
) -> CommodityContextSnapshot:
    snapshot = build_commodity_context_snapshot(results)
    save_commodity_context_snapshot(snapshot, db_path=db_path)
    return snapshot
